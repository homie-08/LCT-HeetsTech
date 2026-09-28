"""Пиктограммы нативными фигурами и схема процесса из фигур."""

from __future__ import annotations

import math
from types import SimpleNamespace

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

from sd.icons import SHAPES, add_icon, icon_paths, pick
from sd.icons.svgpath import arc_to_cubics, bounds, parse_path, parse_svg
from sd.template.model import (Capacity, DesignSystem, Palette, Pattern, SlideSize, Slot,
                               Source, TextStyle, Typography)

W, H = 12192000, 6858000


# --- контуры SVG → кривые -------------------------------------------------------

def test_path_commands_become_absolute_lines_and_curves():
    paths = parse_path("M-5-8h6l4 4v12h-10z")
    assert len(paths) == 1
    kinds = [command[0] for command in paths[0]]
    assert kinds == ["M", "L", "L", "L", "L", "Z"]
    assert paths[0][1] == ("L", 1.0, -8.0)             # h6 от (-5,-8)
    assert paths[0][2] == ("L", 5.0, -4.0)             # l4 4
    assert paths[0][3] == ("L", 5.0, 8.0)              # v12


def test_relative_arc_and_smooth_curves_parse():
    # «money»: круг и знак рубля с дугой; «wave»: гладкое продолжение s.
    for name in ("money", "wave", "team", "lock"):
        paths = icon_paths(name)
        assert paths
        for path in paths:
            assert path[0][0] == "M"
            assert all(command[0] in ("M", "L", "C", "Z") for command in path)


def test_arc_approximation_ends_exactly_at_the_target_and_stays_on_the_circle():
    commands = arc_to_cubics(0, -5, 5, 5, 0, 0, 1, 5, 0)      # четверть окружности r=5
    assert commands[-1][5:] == (5, 0)
    # Середина кривой Безье лежит вблизи окружности радиуса 5.
    _, x1, y1, x2, y2, x, y = commands[0]
    mid = (0.125 * 0 + 0.375 * x1 + 0.375 * x2 + 0.125 * x,
           0.125 * -5 + 0.375 * y1 + 0.375 * y2 + 0.125 * y)
    assert abs(math.hypot(*mid) - 5) < 0.02


def test_every_icon_fits_the_set_box():
    for name in SHAPES:
        low_x, low_y, high_x, high_y = bounds([list(path) for path in icon_paths(name)])
        assert min(low_x, low_y) >= -10.5 and max(high_x, high_y) <= 10.5, name


def test_circle_and_rounded_rect_elements():
    paths = parse_svg('<circle cx="0" cy="0" r="8"/><rect x="-6" y="-1" width="12" height="9" rx="2"/>')
    assert len(paths) == 2
    assert paths[0][0] == ("M", 8.0, 0.0) and paths[0][-1] == ("Z",)
    assert sum(1 for command in paths[1] if command[0] == "C") == 4


# --- нативная фигура --------------------------------------------------------------

def _deck():
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Emu(W), Emu(H)
    return presentation, presentation.slides.add_slide(presentation.slide_layouts[6])


def test_icon_is_a_custom_geometry_shape_with_theme_stroke_and_no_fill():
    presentation, slide = _deck()
    shape = add_icon(slide, "clock", 100000, 200000, 500000, ("scheme", "accent1"))
    element = shape._element
    assert element.find(qn("p:spPr")).find(qn("a:custGeom")) is not None
    assert element.find(qn("p:spPr")).find(qn("a:prstGeom")) is None
    assert element.find(qn("p:style")) is None         # тени и заливки темы сняты
    assert element.find(qn("p:txBody")) is None
    line = element.find(qn("p:spPr")).find(qn("a:ln"))
    assert line.get("cap") == "rnd"
    assert line.find(qn("a:solidFill")).find(qn("a:schemeClr")).get("val") == "accent1"
    assert element.find(qn("p:spPr")).find(qn("a:noFill")) is not None
    paths = element.findall(".//" + qn("a:path"))
    assert len(paths) == 2 and all(path.get("fill") == "none" for path in paths)
    # Файл сохраняется и открывается.
    import io
    buffer = io.BytesIO()
    presentation.save(buffer)
    buffer.seek(0)
    assert len(Presentation(buffer).slides[0].shapes) == 1


def test_icon_takes_an_rgb_color_when_the_role_is_not_a_theme_slot():
    _, slide = _deck()
    shape = add_icon(slide, "mail", 0, 0, 400000, ("rgb", "#70C3FF"))
    line = shape._element.find(qn("p:spPr")).find(qn("a:ln"))
    assert line.find(qn("a:solidFill")).find(qn("a:srgbClr")).get("val") == "70C3FF"


# --- схема процесса ----------------------------------------------------------------

def _design(page: str = "#FFFFFF", accent_slot: str = "theme.accent1") -> DesignSystem:
    return DesignSystem(
        source=Source(file="t.pptx", sha256="0" * 64,
                      slide_size=SlideSize(w_emu=W, h_emu=H, ratio="16:9")),
        palette=Palette(theme={"dk1": "#000000", "lt1": "#FFFFFF", "accent1": "#0077FF"},
                        roles={"page_bg": page, "text_primary": "#000000",
                               "text_inverse": "#FFFFFF", "accent_primary": "#0077FF"},
                        role_provenance={"page_bg": "theme.lt1", "text_primary": "theme.dk1",
                                         "text_inverse": "usage.text_on_dark",
                                         "accent_primary": accent_slot}),
        typography=Typography(scale_pt=[36, 16], styles={"title": TextStyle(size_pt=36),
                                                         "body_l1": TextStyle(size_pt=16)}),
    )


def test_process_diagram_draws_cards_badges_icons_and_arrows():
    from sd.compose.diagram import add_process

    _, slide = _deck()
    steps = ["Заявка поступает из почты или портала",
             "Модель распознаёт тип документа и реквизиты",
             "Правила маршрутизации назначают исполнителя",
             "Система контролирует срок и эскалирует просрочку"]
    notes: list[str] = []
    add_process(slide, [0.05, 0.22, 0.9, 0.6], steps, _design(), W, H, notes, number=3)
    names = [shape.name for shape in slide.shapes]
    assert sum(name.endswith("карточка") for name in names) == 4
    assert sum(name.endswith("номер") for name in names) == 4
    assert sum(name.endswith("текст") for name in names) == 4
    assert sum(name.startswith("Стрелка") for name in names) == 3
    assert any(name.startswith("Пиктограмма") for name in names)     # «почты» → mail
    assert notes == ["слайд 3: схема процесса из фигур — 4 шагов"]
    texts = [shape.text_frame.text for shape in slide.shapes if shape.name.endswith("текст")]
    assert "Заявка поступает из почты или портала" in texts[0]
    # Все фигуры внутри отведённого поля.
    for shape in slide.shapes:
        assert shape.left >= int(0.05 * W) - 1 and shape.left + shape.width <= int(0.95 * W) + 1
        assert shape.top >= int(0.22 * H) - 1 and shape.top + shape.height <= int(0.82 * H) + 1


def test_five_steps_go_in_two_rows_as_a_snake():
    from sd.compose.diagram import add_process

    _, slide = _deck()
    steps = [f"Шаг {n}: действие номер {n}" for n in range(1, 6)]
    add_process(slide, [0.05, 0.22, 0.9, 0.6], steps, _design(), W, H, [], number=1)
    cards = [shape for shape in slide.shapes if shape.name.endswith("карточка")]
    tops = sorted({shape.top for shape in cards})
    assert len(tops) == 2                                   # два ряда
    row_two = [shape for shape in cards if shape.top == tops[1]]
    # Второй ряд идёт справа налево: карточка 4 правее карточки 5.
    by_name = {shape.name: shape for shape in row_two}
    assert by_name["Шаг 4: карточка"].left > by_name["Шаг 5: карточка"].left
    arrows = [shape.name for shape in slide.shapes if shape.name.startswith("Стрелка")]
    assert len(arrows) == 4


def test_cards_are_the_accent_with_transparency_and_text_follows_the_page():
    from sd.compose.diagram import add_process

    _, slide = _deck()
    add_process(slide, [0.05, 0.22, 0.9, 0.6], ["а б", "в г", "д е"],
                _design(page="#101010"), W, H, [], number=1)
    card = next(shape for shape in slide.shapes if shape.name.endswith("карточка"))
    fill = card._element.find(qn("p:spPr")).find(qn("a:solidFill")).find(qn("a:schemeClr"))
    assert fill.get("val") == "accent1"                  # цвет шаблона, а не новый оттенок
    assert fill.find(qn("a:alpha")) is not None
    text = next(shape for shape in slide.shapes if shape.name.endswith("текст"))
    run = text.text_frame.paragraphs[0].runs[0]
    # Белый текст на тёмной странице — слотом темы lt1: значение совпало со слотом.
    assert run.font.color.theme_color.name == "BACKGROUND_1"


def test_role_color_prefers_a_theme_slot_with_the_same_value():
    from sd.compose.diagram import _role_color

    design = _design(accent_slot="usage.accent")          # роль выведена из употребления
    assert _role_color(design, "accent_primary") == ("scheme", "accent1")


# --- синтетический паттерн и матчер --------------------------------------------------

def test_synthetic_pattern_is_built_on_the_calmest_layout_with_a_title():
    from sd.template.synthetic import process_pattern

    design = _design()
    title = Slot(role="title", bbox=[0.06, 0.08, 0.88, 0.1], capacity=Capacity(chars=60))
    design.patterns = [
        Pattern(id="section-l1", archetype="section", source_kind="layout", source_name="Раздел",
                source_part="ppt/slideLayouts/slideLayout1.xml", render_mode="use_layout",
                slots=[title.model_copy()]),
        Pattern(id="metrics-l2", archetype="metrics", source_kind="layout", source_name="Цифры",
                source_part="ppt/slideLayouts/slideLayout2.xml", render_mode="use_layout",
                slots=[title.model_copy(),
                       Slot(role="metric_value", bbox=[0.1, 0.3, 0.2, 0.2], capacity=Capacity(chars=6))]),
        Pattern(id="bullets-s3", archetype="bullets", source_kind="slide", source_name="Пример",
                source_part="ppt/slides/slide3.xml", render_mode="clone_slide",
                slots=[title.model_copy(), Slot(role="body", bbox=[0.06, 0.25, 0.88, 0.6],
                                                 capacity=Capacity(chars=400, lines=8))]),
    ]
    pattern = process_pattern(design)
    assert pattern is not None and pattern.id == "process-diagram"
    assert pattern.source_name == "Раздел" and pattern.archetype == "process"
    assert [slot.role for slot in pattern.slots] == ["title", "diagram"]
    diagram = pattern.slots[1]
    assert diagram.bbox[2] >= 0.55 and diagram.bbox[3] >= 0.34
    assert pattern.accepts["synthetic"] == "process"


def test_matcher_gives_the_diagram_to_a_process_of_three_to_six_steps_only():
    from sd.content.model import ContentIR
    from sd.plan.match import score
    from sd.plan.model import SlidePlan
    from sd.template.synthetic import process_pattern

    design = _design()
    title = Slot(role="title", bbox=[0.06, 0.08, 0.88, 0.1], capacity=Capacity(chars=60))
    design.patterns = [Pattern(id="section-l1", archetype="section", source_kind="layout",
                               source_name="Раздел", source_part="ppt/slideLayouts/slideLayout1.xml",
                               render_mode="use_layout", slots=[title])]
    diagram = process_pattern(design)
    ir = ContentIR()
    four = SlidePlan(n=2, intent="process", key_message="Как это работает",
                     items=[f"шаг {n}" for n in range(4)])
    assert score(diagram, four, ir, 0, design).total > 2.5
    assert score(diagram, four, ir, 0, design).handicap > 0

    two = four.model_copy(update={"items": ["шаг 1", "шаг 2"]})
    assert score(diagram, two, ir, 0, design).total < -100                 # не схема
    seven = four.model_copy(update={"items": [f"шаг {n}" for n in range(7)]})
    assert score(diagram, seven, ir, 0, design).total < -100
    bullets = four.model_copy(update={"intent": "bullets"})
    assert score(diagram, bullets, ir, 0, design).total < score(diagram, four, ir, 0, design).total


def test_fitter_hands_all_steps_to_the_diagram_slot():
    from sd.content.model import ContentIR
    from sd.fit.layout import fill_slide
    from sd.plan.model import ScoreBreakdown, SlidePlan, SlideSpec
    from sd.template.synthetic import process_pattern

    design = _design()
    title = Slot(role="title", bbox=[0.06, 0.08, 0.88, 0.1], capacity=Capacity(chars=60))
    design.patterns = [Pattern(id="section-l1", archetype="section", source_kind="layout",
                               source_name="Раздел", source_part="ppt/slideLayouts/slideLayout1.xml",
                               render_mode="use_layout", slots=[title])]
    pattern = process_pattern(design)
    plan = SlidePlan(n=2, intent="process", key_message="Как это работает",
                     items=["Приём", "Распознавание", "Маршрутизация", "Контроль"])
    filled = fill_slide(design, pattern, SlideSpec(n=2, plan=plan, pattern_id=pattern.id,
                                                   score=ScoreBreakdown()), ContentIR())
    diagram = next(fill for fill in filled.fills if fill.slot.role == "diagram")
    assert diagram.kind == "diagram" and diagram.items == plan.items
    assert filled.dropped == []
    assert next(fill for fill in filled.fills if fill.slot.role == "title").text == "Как это работает"


@pytest.mark.skipif(
    not (__import__("pathlib").Path(__file__).resolve().parents[2] / "ресурсы" / "Датасет").exists(),
    reason="нет датасета организаторов")
def test_vk_templates_get_a_diagram_pattern_and_render_it_natively(tmp_path):
    from pathlib import Path

    from sd.content import parse_content
    from sd.fit.layout import fill_deck
    from sd.compose import build_deck
    from sd.plan import build_plan, match_deck
    from sd.template.analyze import analyze_template

    root = Path(__file__).resolve().parents[2]
    template = root / "ресурсы" / "Датасет" / "VK Tech шаблон.pptx"
    design = analyze_template(template)
    assert design.pattern("process-diagram") is not None
    ir = parse_content(root / "data" / "content" / "product")
    plan = build_plan(ir, None)
    steps = next(slide for slide in plan.slides if slide.intent == "process" and slide.items)
    steps.items = ["Документ поступает из почты", "Модель распознаёт реквизиты",
                   "Правила назначают исполнителя", "Система контролирует срок"]
    spec = match_deck(plan, design, ir)
    target = next(item for item in spec.slides if item.n == steps.n)
    target.pattern_id = "process-diagram"
    spec.slides = [target]
    filled = fill_deck(design, spec, ir)
    result = build_deck(design, filled, ir, template, tmp_path / "deck.pptx")
    assert any("схема процесса" in note for note in result.notes)
    slide = Presentation(str(tmp_path / "deck.pptx")).slides[0]
    names = [shape.name for shape in slide.shapes]
    assert sum(name.endswith("карточка") for name in names) == 4
    assert not any(shape.shape_type == 13 for shape in slide.shapes)      # ни одной картинки


@pytest.mark.skipif(
    not (__import__("pathlib").Path(__file__).resolve().parents[2] / "ресурсы" / "Датасет").exists(),
    reason="нет датасета организаторов")
def test_single_empty_photo_slot_gets_a_pictogram_and_the_fill_knows_it(tmp_path):
    from pathlib import Path

    from sd.content import parse_content
    from sd.fit.layout import fill_deck
    from sd.compose import build_deck
    from sd.plan import build_plan, match_deck
    from sd.template.analyze import analyze_template

    root = Path(__file__).resolve().parents[2]
    template = root / "ресурсы" / "Датасет" / "Шаблон презентации VK Education.pptx"
    design = analyze_template(template)
    ir = parse_content(root / "data" / "content" / "product")
    spec = match_deck(build_plan(ir, None), design, ir)
    target = next(item for item in spec.slides if item.plan.intent in ("bullets", "process")
                  and item.plan.items)
    target.pattern_id = "image_text-slideLayout5"        # текст + фотослот без фото
    spec.slides = [target]
    filled = fill_deck(design, spec, ir)
    result = build_deck(design, filled, ir, template, tmp_path / "deck.pptx")

    assert any("пустой фотослот занят пиктограммой" in note for note in result.notes)
    image_fill = next(fill for fill in filled[0].fills if fill.slot.role == "image")
    assert image_fill.kind == "icon" and image_fill.text
    slide = Presentation(str(tmp_path / "deck.pptx")).slides[0]
    icons = [shape for shape in slide.shapes if shape.name.startswith("Пиктограмма")]
    assert len(icons) == 1
    assert icons[0]._element.find(qn("p:spPr")).find(qn("a:custGeom")) is not None
