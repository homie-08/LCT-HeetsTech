"""Экспорт собранной презентации в .pdf и .html.

`.pptx` — рабочий формат, его правят. PDF нужен, чтобы отправить: тот же
движок, что рендерит слайды, сохраняет файл целиком — PowerPoint через COM,
иначе LibreOffice. HTML — чтобы открыть по ссылке без Office: страница с
кадрами слайдов и их текстом под каждым кадром, самодостаточная (картинки
внутри), листается клавишами. Вёрстку из pptx в HTML не переводим: кадр
точнее любого пересказа фигур в CSS.
"""

from __future__ import annotations

import base64
import html
import os
import subprocess
import sys
import tempfile
from pathlib import Path

from .libreoffice import find_soffice

PP_SAVE_AS_PDF = 32


def to_pdf(deck: str | Path, out: str | Path) -> Path:
    """Сохраняет презентацию в PDF первым доступным движком."""
    deck, out = Path(deck).resolve(), Path(out).resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    errors: list[str] = []
    for exporter in (_pdf_powerpoint, _pdf_libreoffice):
        try:
            exporter(deck, out)
            if out.exists() and out.stat().st_size > 0:
                return out
        except Exception as error:                     # noqa: BLE001 — пробуем следующий
            errors.append(f"{exporter.__name__}: {error}")
    raise RuntimeError("PDF не собран: " + "; ".join(errors))


def _pdf_powerpoint(deck: Path, out: Path) -> None:
    if sys.platform != "win32":
        raise RuntimeError("не Windows")
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    app = presentation = None
    borrowed = False
    try:
        try:
            app, borrowed = win32com.client.GetActiveObject("PowerPoint.Application"), True
        except Exception:
            app = win32com.client.DispatchEx("PowerPoint.Application")
        presentation = app.Presentations.Open(str(deck), ReadOnly=True,
                                              Untitled=False, WithWindow=False)
        if out.exists():
            out.unlink()
        presentation.SaveAs(str(out), PP_SAVE_AS_PDF)
    finally:
        if presentation is not None:
            try:
                presentation.Close()
            except Exception:
                pass
        if app is not None and not borrowed:
            try:
                app.Quit()
            except Exception:
                pass
        app = presentation = None
        pythoncom.CoUninitialize()


def _pdf_libreoffice(deck: Path, out: Path, timeout: int = 180) -> None:
    soffice = find_soffice()
    if soffice is None:
        raise RuntimeError("LibreOffice не найден")
    with tempfile.TemporaryDirectory() as work:
        profile = Path(work) / "profile"
        subprocess.run(
            [soffice, "--headless", "--norestore",
             f"-env:UserInstallation=file:///{profile.as_posix()}",
             "--convert-to", "pdf", "--outdir", work, str(deck)],
            check=True, timeout=timeout, capture_output=True,
            env={**os.environ, "SAL_USE_VCLPLUGIN": "svp"},
        )
        pdfs = list(Path(work).glob("*.pdf"))
        if not pdfs:
            raise RuntimeError("LibreOffice не отдал PDF")
        out.write_bytes(pdfs[0].read_bytes())


def to_html(deck: str | Path, images: list[str | Path], out: str | Path,
            title: str = "Презентация") -> Path:
    """Самодостаточная страница: кадры слайдов и текст каждого слайда.

    Текст дублируется под кадром не для красоты — чтобы страница читалась
    поиском и экранным диктором и чтобы его можно было скопировать.
    """
    from pptx import Presentation

    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    texts: list[list[str]] = []
    for slide in Presentation(str(deck)).slides:
        lines = [shape.text_frame.text.strip() for shape in slide.shapes
                 if shape.has_text_frame and shape.text_frame.text.strip()]
        texts.append(lines)

    sections = []
    for index, image in enumerate(images, start=1):
        path = Path(image)
        src = ("data:image/png;base64,"
               + base64.b64encode(path.read_bytes()).decode("ascii")) if path.exists() else ""
        lines = texts[index - 1] if index - 1 < len(texts) else []
        body = "".join(f"<p>{html.escape(line)}</p>" for line in lines)
        picture = (f'<img src="{src}" alt="Слайд {index}">' if src
                   else f'<div class="missing">Слайд {index}: кадр не отрисован</div>')
        sections.append(
            f'<section id="s{index}" aria-label="Слайд {index}">{picture}'
            f'<div class="text">{body}</div></section>')

    page = f"""<!doctype html>
<html lang="ru"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html.escape(title)}</title>
<style>
  body {{ margin: 0; background: #111; color: #ddd; font: 15px/1.45 system-ui, sans-serif; }}
  header {{ padding: 12px 20px; font-weight: 600; }}
  section {{ max-width: 1100px; margin: 0 auto 32px; padding: 0 16px; }}
  img {{ width: 100%; height: auto; display: block; border-radius: 6px; box-shadow: 0 2px 16px #0008; }}
  .text {{ padding: 8px 4px; color: #aaa; }}
  .text p {{ margin: 2px 0; }}
  .missing {{ aspect-ratio: 16/9; display: grid; place-items: center; background: #222; }}
  @media print {{ body {{ background: #fff; }} .text {{ display: none; }} section {{ break-after: page; }} }}
</style></head>
<body><header>{html.escape(title)} · {len(images)} слайдов</header>
{"".join(sections)}
<script>
  // Стрелки листают по слайду, как в докладе.
  document.addEventListener("keydown", (event) => {{
    const sections = [...document.querySelectorAll("section")];
    const current = sections.findIndex((s) => s.getBoundingClientRect().bottom > 80);
    const step = event.key === "ArrowRight" || event.key === "PageDown" ? 1
               : event.key === "ArrowLeft" || event.key === "PageUp" ? -1 : 0;
    if (!step) return;
    const next = sections[Math.min(sections.length - 1, Math.max(0, current + step))];
    if (next) {{ event.preventDefault(); next.scrollIntoView({{ behavior: "smooth" }}); }}
  }});
</script></body></html>"""
    out.write_text(page, encoding="utf-8")
    return out
