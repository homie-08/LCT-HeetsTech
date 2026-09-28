"""Превью макетов, слайдов-примеров и паттернов.

Ключевой приём — «зонд макетов»: из шаблона собирается временная презентация,
в которой на каждый макет приходится один **пустой** слайд. Её рендер показывает
ровно то, что шаблон рисует сам: плашки, логотипы, фоновые изображения. По этой
картинке дальше находится безопасная зона — то, чего нет ни в одном XML.
"""

from __future__ import annotations

import io
from pathlib import Path

from pptx import Presentation

from ..ooxml.clone import drop_all_slides
from ..ooxml.package import Package
from ..template.model import DesignSystem
from . import render_deck
from .base import CACHE_ROOT

PROBE_DIR = CACHE_ROOT.parent / "probe"


def layout_probe_deck(pkg: Package) -> tuple[Path, list[str]]:
    """Презентация из пустых слайдов — по одному на каждый макет.

    Файл переиспользуется: кеш рендера ключуется по содержимому зонда, а
    python-pptx каждый раз пишет байт-в-байт разный архив (в zip попадают
    отметки времени). Пересобирали бы — кеш не попадал бы никогда.
    """
    PROBE_DIR.mkdir(parents=True, exist_ok=True)
    target = PROBE_DIR / f"{pkg.sha256[:16]}-layouts.pptx"
    if target.exists():
        return target, list(pkg.layouts)

    presentation = Presentation(io.BytesIO(pkg.as_pptx_bytes()))
    drop_all_slides(presentation)

    order: list[str] = []
    for master in presentation.slide_masters:
        for layout in master.slide_layouts:
            slide = presentation.slides.add_slide(layout)
            # `add_slide` копирует плейсхолдеры макета, а пустой плейсхолдер в
            # некоторых шаблонах рисуется сплошной плашкой и закрывает декор.
            # Зонд должен показывать только то, что шаблон рисует сам.
            for placeholder in list(slide.placeholders):
                placeholder._element.getparent().remove(placeholder._element)
            order.append(str(layout.part.partname).lstrip("/"))

    presentation.save(target)
    return target, order


def deck_copy(pkg: Package) -> Path:
    """Копия шаблона как обычной презентации — для рендера слайдов-примеров."""
    PROBE_DIR.mkdir(parents=True, exist_ok=True)
    target = PROBE_DIR / f"{pkg.sha256[:16]}-slides.pptx"
    if not target.exists():
        target.write_bytes(pkg.as_pptx_bytes())
    return target


def render_layouts(pkg: Package, out_dir: Path, width_px: int = 1280,
                   prefer: str | None = None) -> dict[str, Path]:
    """Пустой рендер каждого макета: {часть макета -> png}."""
    deck, order = layout_probe_deck(pkg)
    result = render_deck(deck, out_dir / "layouts", width_px, prefer)
    return dict(zip(order, result.images))


def render_slides(pkg: Package, out_dir: Path, width_px: int = 1280,
                  prefer: str | None = None) -> dict[str, Path]:
    """Рендер слайдов-примеров шаблона: {часть слайда -> png}."""
    if not pkg.slides:
        return {}
    result = render_deck(deck_copy(pkg), out_dir / "slides", width_px, prefer)
    return dict(zip(pkg.slides, result.images))


def render_patterns(design: DesignSystem, pkg: Package, out_dir: Path,
                    width_px: int = 1280, prefer: str | None = None) -> dict[str, Path]:
    """Превью для каждого паттерна — берём картинку его источника."""
    sources: dict[str, Path] = {}
    sources.update(render_layouts(pkg, out_dir, width_px, prefer))
    sources.update(render_slides(pkg, out_dir, width_px, prefer))

    previews: dict[str, Path] = {}
    for pattern in design.patterns:
        image = sources.get(pattern.source_part)
        if image is None:
            continue
        previews[pattern.id] = image
        pattern.preview = str(image)
    return previews
