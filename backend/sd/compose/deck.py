"""Сборка `.pptx`.

Выходной файл — это **копия шаблона**, из которой удалены слайды-примеры. Так
наследуются тема, мастера, макеты, встроенные шрифты и размер слайда: ничего из
оформления воспроизводить руками не нужно, а значит нечему и разойтись с
оригиналом.

Дальше на каждый слайд работает один из двух режимов. `use_layout` добавляет
слайд на штатном макете и заполняет плейсхолдеры — форматирование приходит из
шаблона само. `clone_slide` копирует слайд-пример целиком и подменяет в нём
содержимое: это единственный способ сохранить дизайн шаблонов, где макеты
пустые, а вся вёрстка нарисована руками.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Emu

from ..content.model import ContentIR
from ..fit.layout import FilledSlide, SlotFill
from ..ooxml.clone import (clear_shapes, copy_slide_background, copy_slide_shapes,
                           drop_all_slides, find_shape)
from ..ooxml.package import Package
from ..ooxml.text import clear_autofit_scale, set_text
from ..template.model import DesignSystem, Pattern, Slot

CHART_TYPES = {
    "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
}

PH_TYPE_NAMES = {
    "title": "TITLE", "ctrTitle": "CENTER_TITLE", "subTitle": "SUBTITLE",
    "body": "BODY", "obj": "OBJECT", "pic": "PICTURE", "chart": "CHART",
    "tbl": "TABLE",
}


@dataclass
class BuildResult:
    path: Path
    slides: int = 0
    mode_counts: dict[str, int] = field(default_factory=dict)
    defects: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


# --- вспомогательное ---------------------------------------------------------

def _box(slot: Slot, width: int, height: int) -> tuple[Emu, Emu, Emu, Emu]:
    x, y, w, h = slot.bbox
    return Emu(int(x * width)), Emu(int(y * height)), Emu(int(w * width)), Emu(int(h * height))


def _layout_index(presentation: Presentation) -> dict[str, object]:
    index: dict[str, object] = {}
    for master in presentation.slide_masters:
        for layout in master.slide_layouts:
            index[str(layout.part.partname).lstrip("/")] = layout
    return index


def _find_placeholder(slide, slot: Slot):
    if slot.ph_idx is not None:
        for placeholder in slide.placeholders:
            if placeholder.placeholder_format.idx == slot.ph_idx:
                return placeholder
    wanted = PH_TYPE_NAMES.get(slot.ph_type or "")
    if wanted:
        for placeholder in slide.placeholders:
            if str(placeholder.placeholder_format.type).split(" ")[0] == wanted:
                return placeholder
    return None


def _lines_of(fill: SlotFill) -> list[str]:
    if fill.kind != "items":
        return [fill.text]
    return _tidy_items(fill.items)


def _tidy_items(items: list[str]) -> list[str]:
    """Пункты списка — это тезисы, а не предложения.

    Короткий пункт с точкой на конце читается как обрывок абзаца; в презентациях
    её не ставят. Длинные пункты не трогаем — там точка осмысленна.
    """
    if not items or any(len(item) > 90 for item in items):
        return items
    return [item.rstrip().rstrip(".") if item.rstrip().endswith(".")
            and item.count(".") == 1 else item for item in items]


def _write_text(shape, fill: SlotFill, width: int = 0, height: int = 0) -> None:
    tx_body = shape.text_frame._txBody
    clear_autofit_scale(tx_body)
    set_text(tx_body, _lines_of(fill), fill.size_pt, fill.lead_size_pt)
    if width and height:
        _apply_geometry(shape, fill.slot, width, height)


def _apply_geometry(shape, slot: Slot, width: int, height: int) -> None:
    """Фиксирует положение слота на слайде.

    Плейсхолдер наследует геометрию от макета, но после переразметки она уже
    другая. Заодно это выравнивает то, что видит нормоконтроль: он проверяет
    контент по границам слотов, и они должны совпадать с реальными.
    """
    left, top, box_w, box_h = _box(slot, width, height)
    shape.left, shape.top, shape.width, shape.height = left, top, box_w, box_h


# --- вставка медиа -----------------------------------------------------------

def _compatible_series(series: dict[str, list[float]], ratio: float = 25.0
                       ) -> tuple[dict[str, list[float]], list[str]]:
    """Оставляет ряды сопоставимого масштаба.

    «Заявок обработано» и «средний срок в днях» на одной оси дают столбики
    12400 и 6 — второй ряд физически не виден. Такие ряды лучше не рисовать
    вовсе, чем рисовать невидимыми.
    """
    scales = {name: max((abs(value) for value in values), default=0.0)
              for name, values in series.items()}
    leader = max(scales.values(), default=0.0)
    if leader <= 0:
        return series, []
    kept = {name: values for name, values in series.items()
            if scales[name] * ratio >= leader}
    return kept, [name for name in series if name not in kept]


def _add_chart(slide, slot: Slot, block, width: int, height: int,
               notes: list[str] | None = None) -> None:
    series, dropped = _compatible_series(block.y)
    if dropped and notes is not None:
        notes.append("ряды другого масштаба не показаны на графике: "
                     + ", ".join(dropped))

    data = CategoryChartData()
    data.categories = block.x
    for name, values in series.items():
        data.add_series(name, values)
    left, top, box_w, box_h = _box(slot, width, height)
    chart_type = CHART_TYPES.get(block.suggest or "bar", XL_CHART_TYPE.COLUMN_CLUSTERED)
    frame = slide.shapes.add_chart(chart_type, left, top, box_w, box_h, data)
    # Заголовок диаграммы — явный, а не автоматический от имени ряда: аудит
    # (и PowerPoint при удалении ряда) считает подписанной только такую.
    frame.chart.has_title = True
    frame.chart.chart_title.text_frame.text = (
        block.label or (next(iter(series)) if len(series) == 1 else ""))
    # Легенда нужна только при нескольких рядах: на одном она занимает место зря.
    frame.chart.has_legend = len(series) > 1
    if frame.chart.has_legend:
        frame.chart.legend.include_in_layout = False


def _add_table(slide, slot: Slot, block, width: int, height: int) -> None:
    rows = len(block.rows) + 1
    columns = max(len(block.header), max((len(row) for row in block.rows), default=1))
    left, top, box_w, box_h = _box(slot, width, height)
    table = slide.shapes.add_table(rows, columns, left, top, box_w, box_h).table

    for index, title in enumerate(block.header[:columns]):
        table.cell(0, index).text = title
    for row_index, row in enumerate(block.rows, start=1):
        for column_index, value in enumerate(row[:columns]):
            table.cell(row_index, column_index).text = value

    _fit_table(table, rows, box_h)
    _spread_columns(table, block, columns, box_w)


# Кегль по умолчанию у таблицы крупный, а строки PowerPoint тянет под текст:
# длинная ячейка переносится на две строки, и таблица вылезает за своё поле.
# Считаем кегль от высоты, которая на строку реально приходится.
TABLE_MAX_PT = 18.0
TABLE_MIN_PT = 10.0


def _fit_table(table, rows: int, box_h: int) -> None:
    from pptx.util import Pt

    per_row_pt = box_h / rows / 12700
    # Полторы строки на ячейку — запас на перенос длинного заголовка,
    # плюс поля ячейки сверху и снизу.
    size = max(TABLE_MIN_PT, min(TABLE_MAX_PT, (per_row_pt - 8) / 1.5 / 1.2))
    for row in table.rows:
        for cell in row.cells:
            cell.margin_top = cell.margin_bottom = Pt(3)
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(round(size, 1))


def _spread_columns(table, block, columns: int, box_w: int) -> None:
    """Ширина колонки по длине её текста, а не поровну.

    Поровну — значит колонка с названием получит столько же, сколько колонка
    с «0», и название перенесётся посреди слова. Доля колонки считается по
    самой длинной ячейке в ней, с ограничением снизу, чтобы узкие колонки не
    выродились в полоску.
    """
    from pptx.util import Emu

    rows = [block.header, *block.rows]
    # Плюс два знака на поля ячейки: без запаса заголовок ровно по ширине
    # колонки всё равно переносится.
    longest = [max((len(row[index]) for row in rows if index < len(row)), default=1) + 2
               for index in range(columns)]
    floor = max(longest) * 0.25          # не уже четверти самой широкой
    weights = [max(value, floor) for value in longest]
    total = sum(weights) or 1.0

    used = 0
    for index, weight in enumerate(weights):
        # Последней колонке отдаём остаток: округление не должно менять ширину.
        share = box_w - used if index == columns - 1 else int(box_w * weight / total)
        table.columns[index].width = Emu(max(1, share))
        used += share


def _add_picture(slide, slot: Slot, path: str, width: int, height: int) -> None:
    left, top, box_w, box_h = _box(slot, width, height)
    picture = slide.shapes.add_picture(path, left, top)
    # Вписываем с сохранением пропорций: растянутая фотография заметна сразу.
    scale = min(box_w / picture.width, box_h / picture.height)
    picture.width = int(picture.width * scale)
    picture.height = int(picture.height * scale)
    picture.left = int(left + (box_w - picture.width) / 2)
    picture.top = int(top + (box_h - picture.height) / 2)


def _place_media(slide, fill: SlotFill, ir: ContentIR, width: int, height: int,
                 notes: list[str], design: DesignSystem | None = None, number: int = 0) -> None:
    if fill.kind == "diagram":
        from .diagram import add_process

        if design is not None and fill.items:
            add_process(slide, list(fill.slot.bbox), fill.items, design, width, height,
                        notes, number)
        return
    block = ir.block(fill.block_id) if fill.block_id else None
    if block is None:
        return
    placeholder = _find_placeholder(slide, fill.slot)
    if placeholder is not None:
        placeholder._element.getparent().remove(placeholder._element)

    if fill.kind == "chart":
        _add_chart(slide, fill.slot, block, width, height, notes)
    elif fill.kind == "table":
        _add_table(slide, fill.slot, block, width, height)
    elif fill.kind == "image":
        asset = next((a for a in ir.assets if a.id == block.asset_id), None)
        if asset and Path(asset.path).exists():
            _add_picture(slide, fill.slot, asset.path, width, height)
        else:
            notes.append(f"слайд: изображение {block.asset_id} недоступно")


# --- два режима сборки -------------------------------------------------------

def _compose_from_layout(presentation, layout, filled: FilledSlide, ir: ContentIR,
                         width: int, height: int, notes: list[str],
                         design: DesignSystem | None = None):
    slide = presentation.slides.add_slide(layout)
    used = set()

    for fill in filled.fills:
        if fill.kind in ("chart", "table", "image", "diagram"):
            _place_media(slide, fill, ir, width, height, notes, design, filled.n)
            continue
        if fill.is_empty:
            continue
        placeholder = _find_placeholder(slide, fill.slot)
        if placeholder is None:
            notes.append(f"слайд {filled.n}: плейсхолдер {fill.slot.role} не найден в макете")
            continue
        _write_text(placeholder, fill, width, height)
        used.add(placeholder.placeholder_format.idx)

    # Незаполненный плейсхолдер не бывает невидимым: текстовый рисуется
    # подсказкой «Текст заголовка», а плейсхолдер картинки в некоторых шаблонах —
    # сплошной цветной плашкой. Убираем любые незанятые; на месте пустого
    # фотослота, если у слайда есть текст, встаёт пиктограмма по его смыслу —
    # макет под фото без фото выглядит недоделанным.
    context = " ".join(fill.text + " " + " ".join(fill.items) for fill in filled.fills
                       if fill.kind in ("text", "items"))
    spare = [placeholder for placeholder in slide.placeholders
             if placeholder.placeholder_format.idx not in used
             and not (placeholder.has_text_frame and placeholder.text_frame.text.strip())]
    # Одна пиктограмма на слайд: четыре одинаковых значка в галерее из четырёх
    # рамок — не иллюстрация, а заглушка.
    empty_photos = [placeholder for placeholder in spare if _is_picture_placeholder(placeholder)]
    for placeholder in spare:
        if (design is not None and context.strip() and len(empty_photos) == 1
                and placeholder is empty_photos[0]
                and filled.pattern.archetype in ICON_FOR_EMPTY_PHOTO):
            _icon_instead_of_photo(slide, placeholder, context, design, width, height,
                                   notes, filled)
        placeholder._element.getparent().remove(placeholder._element)
    return slide


# Макеты, где пустой фотослот заменяется пиктограммой по смыслу текста.
ICON_FOR_EMPTY_PHOTO = {"image_text", "gallery", "bullets", "two_column"}
# Доля меньшей стороны фотослота под пиктограмму.
ICON_SHARE = 0.42


def _is_picture_placeholder(placeholder) -> bool:
    from pptx.enum.shapes import PP_PLACEHOLDER

    try:
        # Только плейсхолдер картинки: «объект» — это прежде всего текст, и
        # значок посреди пустого текстового поля был бы посторонним.
        return placeholder.placeholder_format.type == PP_PLACEHOLDER.PICTURE
    except Exception:                                 # noqa: BLE001 — не плейсхолдер
        return False


def _icon_instead_of_photo(slide, placeholder, context: str, design: DesignSystem,
                           width: int, height: int, notes: list[str],
                           filled: FilledSlide) -> None:
    from ..icons import add_icon, pick
    from .diagram import contrasting_color

    box = _shape_box(placeholder, width, height)
    if box is None:
        return
    name = pick(context)
    if name is None:
        return
    size = int(min(box[2] * width, box[3] * height) * ICON_SHARE)
    if size <= 0:
        return
    left = int(box[0] * width + (box[2] * width - size) / 2)
    top = int(box[1] * height + (box[3] * height - size) / 2)
    add_icon(slide, name, left, top, size, contrasting_color(design, slide, "accent_primary"),
             label=f"вместо фото, слайд {filled.n}")
    # Укладка узнаёт о значке: слот больше не пустой, и нормоконтроль по
    # рендеру ждёт содержание именно здесь.
    for fill in filled.fills:
        if fill.slot.role == "image" and fill.is_empty and _mostly_same(fill.slot.bbox, box):
            fill.kind = "icon"
            fill.text = name
            break
    notes.append(f"слайд {filled.n}: пустой фотослот занят пиктограммой «{name}»")


def _mostly_same(a, b, share: float = 0.5) -> bool:
    """Один и тот же бокс с точностью до обрезки по краю слайда: пересечение
    покрывает больше половины меньшего из двух."""
    left, top = max(a[0], b[0]), max(a[1], b[1])
    right = min(a[0] + a[2], b[0] + b[2])
    bottom = min(a[1] + a[3], b[1] + b[3])
    if right <= left or bottom <= top:
        return False
    smaller = min(a[2] * a[3], b[2] * b[3]) or 1e-9
    return (right - left) * (bottom - top) / smaller >= share


def _compose_from_clone(presentation, layout, pkg: Package, pattern: Pattern,
                        filled: FilledSlide, ir: ContentIR, width: int, height: int,
                        notes: list[str], design: DesignSystem | None = None):
    slide = presentation.slides.add_slide(layout)
    clear_shapes(slide)
    notes.extend(f"слайд {filled.n}: {note}"
                 for note in copy_slide_shapes(pkg, pattern.source_part, slide))
    copy_slide_background(pkg, pattern.source_part, slide)

    for fill in filled.fills:
        if fill.kind in ("chart", "table", "image", "diagram"):
            if fill.slot.shape_id:
                if (shape := find_shape(slide, fill.slot.shape_id)) is not None:
                    shape._element.getparent().remove(shape._element)
            _place_media(slide, fill, ir, width, height, notes, design, filled.n)
            continue
        if fill.is_empty or not fill.slot.shape_id:
            continue
        shape = find_shape(slide, fill.slot.shape_id)
        if shape is None or not shape.has_text_frame:
            notes.append(f"слайд {filled.n}: шейп {fill.slot.shape_id} не найден в копии")
            continue
        _write_text(shape, fill, width, height)

    # Текст, который не заменили, остался из примера — его нужно убрать.
    filled_ids = {fill.slot.shape_id for fill in filled.fills if not fill.is_empty}
    empty_slots = [slot for slot in pattern.slots
                   if slot.shape_id and slot.shape_id not in filled_ids]

    # Сначала карточки целиком: плашка и её акцентная полоска без текста
    # выглядят как забытая пустая ячейка, а не как замысел дизайнера.
    _drop_empty_cards(slide, pattern, empty_slots, filled_ids, width, height, notes,
                      filled.n)

    for slot in empty_slots:
        shape = find_shape(slide, slot.shape_id)
        if shape is not None and shape.has_text_frame:
            shape._element.getparent().remove(shape._element)
    return slide


# Шейп крупнее этой доли слайда — фон или рамка, а не карточка.
CARD_MAX_AREA = 0.55
# Допуск на вложенность: акцентная полоска часто лежит ровно по краю плашки.
CARD_MARGIN = 0.012


def _shape_box(shape, width: int, height: int) -> tuple[float, float, float, float] | None:
    """Габариты шейпа в долях слайда."""
    if None in (shape.left, shape.top, shape.width, shape.height):
        return None
    return (shape.left / width, shape.top / height,
            shape.width / width, shape.height / height)


def _contains(outer: tuple[float, float, float, float],
              inner: tuple[float, float, float, float], margin: float = 0.0) -> bool:
    return (outer[0] - margin <= inner[0]
            and outer[1] - margin <= inner[1]
            and outer[0] + outer[2] + margin >= inner[0] + inner[2]
            and outer[1] + outer[3] + margin >= inner[1] + inner[3])


def _drop_empty_cards(slide, pattern: Pattern, empty_slots, filled_ids,
                      width: int, height: int, notes: list[str], number: int) -> None:
    """Удаляет плашки, внутри которых не осталось ни одной заполненной надписи.

    Карточка в разметке — это не одна фигура, а плашка, полоска-акцент и
    надпись. Убрав только надпись, мы оставляем пустую коробку. Поэтому ищем
    наименьшую фигуру, охватывающую осиротевшую надпись, и убираем всё, что
    лежит внутри неё. Карточку, где хоть что-то заполнено, не трогаем.
    """
    filled_boxes = [slot.bbox for slot in pattern.slots
                    if slot.shape_id in filled_ids]

    for slot in empty_slots:
        shapes = [(shape, box) for shape in slide.shapes
                  if (box := _shape_box(shape, width, height)) is not None]
        # Сама надпись под условие вложенности тоже подходит — её отсекаем по
        # площади, иначе «наименьшим контейнером» всегда оказывается она.
        slot_area = slot.bbox[2] * slot.bbox[3]
        containers = [
            (shape, box) for shape, box in shapes
            if slot_area * 1.2 < box[2] * box[3] <= CARD_MAX_AREA
            and _contains(box, slot.bbox, CARD_MARGIN)
            and not any(_contains(box, filled, CARD_MARGIN) for filled in filled_boxes)
        ]
        if not containers:
            continue

        card, card_box = min(containers, key=lambda item: item[1][2] * item[1][3])

        removed = 0
        for shape, box in shapes:
            if _contains(card_box, box, CARD_MARGIN):
                parent = shape._element.getparent()
                if parent is not None:
                    parent.remove(shape._element)
                    removed += 1
        if removed:
            notes.append(f"слайд {number}: пустая карточка убрана "
                         f"({removed} фигур)")

    _drop_empty_rows(slide, pattern, empty_slots, filled_boxes, width, height,
                     notes, number)
    _drop_empty_columns(slide, empty_slots, filled_boxes, width, height, notes, number)


# Мелкий спутник надписи: иконка, стрелка, разделитель. Крупнее — это уже
# самостоятельная часть слайда, и трогать её нельзя.
SATELLITE_MAX_AREA = 0.06


def _centre_in_band(box: tuple[float, float, float, float],
                    top: float, bottom: float) -> bool:
    centre = box[1] + box[3] / 2
    return top <= centre <= bottom


def _band_touches(box: tuple[float, float, float, float],
                  top: float, bottom: float) -> bool:
    """Пересекается ли фигура с полосой хотя бы краем."""
    return box[1] < bottom and box[1] + box[3] > top


def _drop_empty_rows(slide, pattern: Pattern, empty_slots, filled_boxes,
                     width: int, height: int, notes: list[str], number: int) -> None:
    """Убирает спутников осиротевшей строки: иконку, стрелку, разделитель.

    Ряд в списке — это не фигура-контейнер, а просто несколько фигур на одной
    высоте. Контейнера у такой строки нет, поэтому ориентируемся на полосу:
    всё мелкое на высоте пустой надписи уходит вместе с ней — но только если в
    этой же полосе не осталось заполненного текста.
    """
    for slot in empty_slots:
        top = slot.bbox[1] - CARD_MARGIN
        bottom = slot.bbox[1] + slot.bbox[3] + CARD_MARGIN
        # Достаточно, чтобы заполненный текст задел полосу краем: у строки
        # обычно два слота — шапка и пояснение, и заполнен может быть любой.
        if any(_band_touches(box, top, bottom) for box in filled_boxes):
            continue

        removed = 0
        for shape in list(slide.shapes):
            box = _shape_box(shape, width, height)
            if box is None or box[2] * box[3] > SATELLITE_MAX_AREA:
                continue
            if not _centre_in_band(box, top, bottom):
                continue
            parent = shape._element.getparent()
            if parent is not None:
                parent.remove(shape._element)
                removed += 1
        if removed:
            notes.append(f"слайд {number}: пустая строка убрана ({removed} фигур)")


# Насколько выше пустой карточки может стоять её иконка (доля высоты слайда).
COLUMN_REACH = 0.3


def _drop_empty_columns(slide, empty_slots, filled_boxes, width: int, height: int,
                        notes: list[str], number: int) -> None:
    """Убирает иконку над осиротевшей карточкой в ряду.

    Карточки одного ряда стоят колонками: иконка сверху, текст под ней. Когда
    тезисов меньше, чем карточек, последняя остаётся без текста — а её иконка
    без строки под ней читается как ошибка вёрстки. Спутника ищем в колонке
    пустого слота, выше него, и только если в этой колонке нет заполненного
    текста.
    """
    for slot in empty_slots:
        left, right = slot.bbox[0], slot.bbox[0] + slot.bbox[2]
        top, bottom = slot.bbox[1] - COLUMN_REACH, slot.bbox[1] + slot.bbox[3] + CARD_MARGIN

        def in_column(box) -> bool:
            centre_x = box[0] + box[2] / 2
            return left <= centre_x <= right and _centre_in_band(box, top, bottom)

        if any(in_column(box) for box in filled_boxes):
            continue
        removed = 0
        for shape in list(slide.shapes):
            box = _shape_box(shape, width, height)
            if box is None or box[2] * box[3] > SATELLITE_MAX_AREA or not in_column(box):
                continue
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                continue                                # чужой текст не трогаем
            parent = shape._element.getparent()
            if parent is not None:
                parent.remove(shape._element)
                removed += 1
        if removed:
            notes.append(f"слайд {number}: иконка пустой карточки убрана ({removed} фигур)")


# --- точка входа -------------------------------------------------------------

def build_deck(design: DesignSystem, filled_slides: list[FilledSlide], ir: ContentIR,
               template: str | Path, out_path: str | Path) -> BuildResult:
    pkg = Package.open(template)
    presentation = Presentation(io.BytesIO(pkg.as_pptx_bytes()))
    drop_all_slides(presentation)

    width, height = pkg.slide_size
    layouts = _layout_index(presentation)
    result = BuildResult(path=Path(out_path))

    for filled in filled_slides:
        pattern = filled.pattern
        layout_part = pattern.layout_part or pattern.source_part
        layout = layouts.get(layout_part) or next(iter(layouts.values()), None)
        if layout is None:
            result.notes.append(f"слайд {filled.n}: в шаблоне нет макетов")
            continue

        if pattern.render_mode == "clone_slide":
            _compose_from_clone(presentation, layout, pkg, pattern, filled, ir,
                                width, height, result.notes, design)
        else:
            _compose_from_layout(presentation, layout, filled, ir, width, height,
                                 result.notes, design)

        result.mode_counts[pattern.render_mode] = \
            result.mode_counts.get(pattern.render_mode, 0) + 1
        result.slides += 1
        result.defects.extend(filled.defects)

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(out_path))
    return result
