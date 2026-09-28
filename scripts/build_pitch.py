"""Сборка питч-презентации проекта на шаблоне ЛЦТ 2026.

Собирается вручную (python-pptx), а не нашим сервисом: это презентация о
сервисе, и её содержание — авторское. От шаблона берутся только его макеты:
ни одного своего цвета, шрифта или координаты здесь нет, текст кладётся в
плейсхолдеры, оформление приходит из шаблона.

    python scripts/build_pitch.py --pdf

Места, которые команда заполняет сама, помечены в тексте квадратными
скобками — их видно и в файле, и в списке `TODO` на выходе скрипта.
"""

from __future__ import annotations

import argparse
import pathlib
import re
import sys

from pptx import Presentation
from pptx.util import Emu

ROOT = pathlib.Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "ресурсы" / "Презентация" / "ЛЦТ2026 Шаблон презентации.pptx"
IMG = ROOT / "data" / "out" / "pitch" / "img"

# Слайды: макет шаблона, текст по индексам плейсхолдеров, картинки.
# Индексы взяты из самого шаблона (см. `--layouts`), чтобы текст ложился
# туда, куда его положил дизайнер.
DECK: list[dict] = [
    {
        "layout": "Титульный слайд",
        "text": {
            0: "Heets Tech",
            12: "Задача 4. VK Tech — цифровой дизайнер презентаций · Томск",
        },
    },
    {
        "layout": "Проблема и решение",
        "title": "Текст пишут за час, вёрстку — целый день",
        "text": {
            33: "Проблема",
            26: "Автор написал содержание — дальше специалист подбирает макет, двигает "
                "блоки, ищет иконку, перекрашивает диаграмму в корпоративный цвет. "
                "Каждую неделю, для каждой презентации.",
            34: "Что уже есть",
            31: "Gamma, Copilot и другие генераторы делают слайды в своих темах. "
                "Чужой шаблон они не разбирают — фирменный стиль остаётся ручной "
                "работой.",
            35: "Наше решение",
            32: "Читаем чужой шаблон как набор правил — дизайн-система, библиотека "
                "макетов, ограничения — и собираем по ним новые слайды. На выходе "
                "нативный .pptx, который правят руками.",
        },
    },
    {
        "layout": "Пункты",
        "title": "Подход: модель отвечает за смысл, код — за пиксели",
        "text": {
            14: "Смысл — модели.\nПиксели — коду.",
            15: "Шаблон разбирается в дизайн-систему и библиотеку паттернов: палитра "
                "с провенансом, шкала кеглей, слоты с ёмкостью по метрикам шрифта.",
            16: "Модель строит план истории и формулировки. Ни один цвет, шрифт или "
                "координата от неё не зависят — иначе результат не проверить.",
            17: "Макет под слайд выбирает венгерский алгоритм по всей колоде сразу, "
                "и каждая оценка раскладывается на слагаемые.",
            18: "Текст укладывается по шкале шаблона: кегль вниз, потом сокращение "
                "моделью, потом обрезка — и только потом «не размещено».",
            19: "Собранный файл проходит чек-лист заказчика; что исправимо — "
                "пересобирается по выбору пользователя.",
        },
    },
    {
        "layout": "Стадии",
        "title": "Архитектура пайплайна: пять слоёв с явными границами",
        "text": {
            26: "Разбор шаблона",
            27: "OOXML и рендер пустых макетов: палитра, типографика, сетка, "
                "паттерны и препятствия, которых нет в XML",
            28: "Генерация",
            29: "Content IR из md, docx, xlsx, csv; план истории моделью, "
                "эвристика — запасной путь",
            30: "Вёрстка",
            31: "Матчер выбирает макет, укладчик считает ёмкость по метрикам "
                "шрифта, сборка даёт нативные объекты",
            32: "Аудит и ремонт",
            33: "30 проверок: 26 детерминированных и 4 контекстуальные; "
                "пользователь выбирает, что чинить",
            34: "Экспорт",
            35: ".pptx, .pdf и .html из одной и той же колоды",
        },
    },
    {
        "layout": "Содержание_1",
        "title": "Что сервис вытаскивает из чужого шаблона",
        "text": {
            49: "01", 37: "Палитра",
            38: "Слоты темы и роли с провенансом: accent_primary #0077FF ← theme.accent1",
            50: "02", 39: "Типографика",
            40: "Пары шрифтов и шкала кеглей; ёмкость слота считается по cmap и hmtx "
                "самого файла шрифта",
            51: "03", 41: "Сетка и поля",
            42: "Направляющие, безопасная зона и отступы — из макетов шаблона, а не "
                "из констант в коде",
            52: "04", 43: "Библиотека паттернов",
            44: "Каждый макет и слайд-пример превращается в слоты с ролями и "
                "ёмкостью; архетип определяется с объяснением",
            53: "05", 45: "Зрение по рендеру",
            46: "Пустые макеты рендерятся: логотипы и плашки находятся там, где их "
                "не видно в разметке",
            54: "06", 47: "Ограничения",
            48: "Чего в шаблоне нет, сервис говорит прямо: под какие архетипы "
                "не нашлось макетов",
        },
    },
    {
        "layout": "1_Демо_1",
        "title": "От запроса до трёх вариантов вёрстки",
        "pictures": {14: "04-структура.png", 18: "02-колода-варианты.png"},
        "text": {
            15: "Структура до вёрстки: состав слайдов, заголовки-выводы и тезисы "
                "видны и правятся до сборки.",
            16: "Три варианта одного контента на одном шаблоне: сжато, "
                "сбалансированно, просторно. Ось различий — плотность.",
        },
    },
    {
        "layout": "Демо_1",
        "title": "Слайд — это объекты, а не картинка",
        "pictures": {14: "06-схема-процесса.png"},
        "text": {
            15: "Схема процесса собирается из фигур в цветах темы: карточки, номера, "
                "стрелки, пиктограммы. Всё правится в PowerPoint.",
            16: "Таблицы и графики вставляются нативными объектами и берут акценты "
                "темы шаблона.",
            17: "Пиктограммы — свой набор из 34 контуров, переведённых в custGeom; "
                "подбираются по смыслу тезиса.",
        },
    },
    {
        "layout": "Заголовок и объект",
        "title": "Один контент, три шаблона организаторов",
        "picture_over": ("05-три-шаблона.png", 1),
        "text": {},
    },
    {
        "layout": "Демо_1",
        "title": "Аудит — часть конвейера, а не внешняя проверка",
        "pictures": {14: "03-аудит.png"},
        "text": {
            15: "Детерминированные проверки: границы и поля, шкала кеглей, палитра, "
                "плотность, заглушки, дубли, цифры не из исходных материалов.",
            16: "Контекстуальные — вопросы модели: вывод ли в заголовке, о том ли "
                "слайд, одна ли мысль, связаны ли соседи. Помечены отдельно.",
            17: "Пользователь отмечает, что чинить: слайд пересобирается ужатым "
                "бюджетом или на макете с бо́льшим числом мест.",
        },
    },
    {
        "layout": "Пункты",
        "title": "Технологии: только открытые веса",
        "text": {
            14: "Открытые веса.\nЗапуск одним файлом.",
            15: "Модель: Qwen3-4B-Instruct, Apache 2.0, через OpenAI-совместимый "
                "endpoint. Переключение на инференс организаторов — три строки в .env.",
            16: "Бэкенд: Python 3.11, FastAPI, python-pptx и lxml, fontTools, "
                "OpenCV и scikit-image для зрения по рендеру.",
            17: "Фронтенд: React 18, TypeScript, Vite, Tailwind. API отделён от "
                "интерфейса — сервис встраивается в портал.",
            18: "Рендер и PDF: PowerPoint COM как эталон, LibreOffice как "
                "переносимый запасной путь.",
            19: "Промпты и конфиги агентов — отдельные версионированные файлы; "
                "версия каждого попадает в отчёт сборки.",
        },
    },
    {
        "layout": "1_Статистика",
        "title": "Что уже работает",
        "text": {
            21: "9 колод для сдачи",
            18: "три шаблона организаторов × три варианта вёрстки на одном контенте",
            22: "30 проверок аудита",
            23: "26 детерминированных и 4 контекстуальные по Приложению 1 ТЗ; "
                "6 из них чинятся пересборкой по выбору пользователя",
            24: "6 минут 37 секунд",
            25: "живой прогон от чужого шаблона до готовой презентации на ноутбуке "
                "с локальной 4B-моделью, без монтажных склеек",
            26: "665 автотестов",
            27: "628 на бэкенде и 37 на интерфейсе, включая линтер, который "
                "запрещает цвета, шрифты и координаты в коде генерации",
        },
    },
    {
        "layout": "Пункты",
        "title": "Что дальше",
        "text": {
            14: "Готово\nк пилоту.",
            15: "Инференс организаторов вместо локальной 4B: те же промпты, "
                "меняются три строки в .env — качество заголовков вырастет сразу.",
            16: "Библиотека фирменных схем: к схеме процесса добавить цикл, "
                "воронку и сравнение — теми же нативными фигурами.",
            17: "Генерация изображений внутри слайда — задача со звёздочкой, "
                "модель до 20B подключается тем же адаптером.",
            18: "Массовый режим: один контент через все шаблоны компании одной "
                "командой — в CLI это уже есть.",
            19: "Встраивание в корпоративный портал: HTTP API готов, интерфейс "
                "подключается к нему как обычный клиент.",
        },
    },
    {
        "layout": "Сравнение",
        "title": "Команда «Heets Tech», Томск",
        "text": {
            1: "О команде",
            2: "Капитан: Туранов Тимур Эшанкулович — backend, frontend и iOS "
               "разработчик, младший научный сотрудник и аспирант ТГУ.\n"
               "Участников: 2. Город: Томск.\n"
               "Роли закрыты внутри команды: разбор OOXML и сборка, интерфейс, "
               "работа с моделью и данными.",
            3: "Уникальность решения",
            4: "Мы не генерируем слайды в своей теме, а читаем чужой шаблон как "
               "набор правил и собираем по ним нативный .pptx.\n"
               "Оформление целиком приходит из шаблона: в коде генерации нет ни "
               "одного цвета, шрифта и координаты — это проверяется тестом.\n"
               "Аудит встроен в конвейер, и пользователь сам выбирает, какие "
               "находки чинить.",
        },
    },
    {
        # Макет «Команда» рассчитан на пять карточек шириной в полтора слова;
        # для двоих берём раскладку «фото сверху, подпись снизу» — и ники с
        # телефонами не рвутся, и портреты на месте.
        "layout": "1_Демо_1",
        "title": "Кто делал",
        "pictures": {14: "10-туранов.jpg", 18: "11-протасова.jpg"},
        "text": {
            15: "Туранов Тимур Эшанкулович — капитан\n"
                "Backend, frontend, iOS разработчик\n"
                "@homiekrip · +7 952 887-41-93\n"
                "Младший научный сотрудник и аспирант ТГУ",
            16: "Протасова Ксения Дмитриевна\n"
                "Backend developer, data scientist, frontend developer\n"
                "@grapesfromtheyard · +7 953 929-44-43\n"
                "Самозанятость",
        },
    },
    {
        "layout": "Заголовок и объект",
        "title": "Выводы",
        "text": {
            1: "Сервис читает чужой шаблон как набор правил и собирает по ним новые "
               "слайды — нативные и редактируемые, а не картинки.\n"
               "Смысл доверен модели с открытыми весами, вёрстка и проверки — "
               "детерминированные и воспроизводимые.\n"
               "Аудит встроен в конвейер: находки с номерами слайдов, ремонт по "
               "выбору пользователя.\n"
               "Три варианта вёрстки, экспорт в .pptx, .pdf и .html, запуск одним "
               "файлом и открытый репозиторий.\n"
               "Демо: живой прогон 6:37 без монтажных склеек.",
        },
    },
]


def drop_example_slides(presentation) -> int:
    """Убирает слайды-примеры шаблона, оставляя мастера и макеты."""
    slides = presentation.slides._sldIdLst
    removed = 0
    for element in list(slides):
        rid = element.get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        presentation.part.drop_rel(rid)
        slides.remove(element)
        removed += 1
    return removed


def layouts_of(presentation) -> dict:
    return {layout.name: layout
            for master in presentation.slide_masters
            for layout in master.slide_layouts}


def fill(placeholder, text: str) -> None:
    """Кладёт текст в плейсхолдер, сохраняя оформление первой строки шаблона."""
    frame = placeholder.text_frame
    lines = text.split("\n")
    frame.text = lines[0]
    for line in lines[1:]:
        paragraph = frame.add_paragraph()
        paragraph.text = line
        paragraph.level = frame.paragraphs[0].level


def build(template: pathlib.Path, out: pathlib.Path) -> list[str]:
    presentation = Presentation(str(template))
    drop_example_slides(presentation)
    layouts = layouts_of(presentation)
    todo: list[str] = []

    for number, spec in enumerate(DECK, start=1):
        layout = layouts.get(spec["layout"])
        if layout is None:
            raise SystemExit(f"в шаблоне нет макета «{spec['layout']}»")
        slide = presentation.slides.add_slide(layout)
        by_idx = {ph.placeholder_format.idx: ph for ph in slide.placeholders}

        if spec.get("title") is not None:
            title = slide.shapes.title or by_idx.get(0)
            if title is not None:
                fill(title, spec["title"])

        for idx, text in spec.get("text", {}).items():
            placeholder = by_idx.get(idx)
            if placeholder is None:
                raise SystemExit(f"слайд {number}: в макете нет плейсхолдера {idx}")
            fill(placeholder, text)
            todo.extend(f"слайд {number}: {mark}" for mark in re.findall(r"\[[^\]]+\]", text))

        with_picture: set[int] = set()
        for idx, name in spec.get("pictures", {}).items():
            placeholder = by_idx.get(idx)
            picture = IMG / name
            if placeholder is None or not picture.exists():
                todo.append(f"слайд {number}: нет картинки {name}")
                continue
            box = (placeholder.left, placeholder.top, placeholder.width, placeholder.height)
            shape = placeholder.insert_picture(str(picture))
            shape.crop_left = shape.crop_right = shape.crop_top = shape.crop_bottom = 0
            native_w, native_h = shape.image.size
            ratio = native_w / native_h
            if box[2] / box[3] > ratio:
                height, width = box[3], int(box[3] * ratio)
            else:
                width, height = box[2], int(box[2] / ratio)
            shape.width, shape.height = width, height
            shape.left = int(box[0] + (box[2] - width) / 2)
            shape.top = int(box[1] + (box[3] - height) / 2)
            with_picture.add(idx)

        # Картинка поверх слота, когда слот не «картиночный»: вписываем по его
        # геометрии с сохранением пропорций, сам слот убираем.
        if "picture_over" in spec:
            name, idx = spec["picture_over"]
            slot = by_idx.get(idx)
            picture = IMG / name
            if slot is not None and picture.exists():
                box = (slot.left, slot.top, slot.width, slot.height)
                slot._element.getparent().remove(slot._element)
                add_fitted(slide, picture, *box)

        # Пустые плейсхолдеры в готовом файле показывают подсказки шаблона —
        # убираем всё, что не заполнили (кроме номера слайда и колонтитула).
        for placeholder in list(slide.placeholders):
            kind = str(placeholder.placeholder_format.type)
            if kind.startswith(("SLIDE_NUMBER", "FOOTER", "DATE")):
                continue
            if placeholder.has_text_frame and placeholder.text_frame.text.strip():
                continue
            if placeholder.placeholder_format.idx in with_picture:
                continue                                  # в него вставлена картинка
            placeholder._element.getparent().remove(placeholder._element)

    out.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(out))
    return todo


def add_fitted(slide, picture: pathlib.Path, left: int, top: int,
               width: int, height: int):
    """Вставляет картинку в прямоугольник, сохраняя пропорции и центрируя."""
    shape = slide.shapes.add_picture(str(picture), Emu(left), Emu(top))
    scale = min(width / shape.width, height / shape.height)
    shape.width = int(shape.width * scale)
    shape.height = int(shape.height * scale)
    shape.left = int(left + (width - shape.width) / 2)
    shape.top = int(top + (height - shape.height) / 2)
    return shape


PP_SAVE_PPTX = 24
PP_SAVE_PDF = 32
MSO_TRUE = -1


def to_pdf(source: pathlib.Path, target: pathlib.Path) -> None:
    """PDF тем же движком, что и показывает слайды, — PowerPoint через COM.

    Заодно вшивает шрифты в .pptx: Montserrat стоит не у всех, а без него
    шаблон поедет на чужой машине.
    """
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    app = presentation = None
    try:
        app = win32com.client.DispatchEx("PowerPoint.Application")
        presentation = app.Presentations.Open(str(source.resolve()), ReadOnly=False,
                                              Untitled=False, WithWindow=False)
        # Вшивание задаётся третьим параметром SaveAs: свойство презентации
        # через позднее связывание не выставляется.
        presentation.SaveAs(str(source.resolve()), PP_SAVE_PPTX, MSO_TRUE)
        if target.exists():
            target.unlink()
        presentation.SaveAs(str(target.resolve()), PP_SAVE_PDF)
    finally:
        if presentation is not None:
            presentation.Close()
        if app is not None:
            app.Quit()
        app = presentation = None
        pythoncom.CoUninitialize()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", default="materials/pitch-heets-tech.pptx")
    parser.add_argument("--pdf", action="store_true", help="сохранить ещё и .pdf")
    parser.add_argument("--layouts", action="store_true",
                        help="показать макеты шаблона и их плейсхолдеры")
    args = parser.parse_args()

    if args.layouts:
        presentation = Presentation(str(TEMPLATE))
        for name, layout in layouts_of(presentation).items():
            print(f"[{name}]")
            for ph in layout.placeholders:
                fmt = ph.placeholder_format
                print(f"    {fmt.idx}: {fmt.type}")
        return 0

    out = pathlib.Path(args.out)
    if not out.is_absolute():
        out = ROOT / out
    todo = build(TEMPLATE, out)
    print(f"собрано: {out} — {len(DECK)} слайдов, {out.stat().st_size / 1e6:.0f} МБ")
    if args.pdf:
        pdf = out.with_suffix(".pdf")
        to_pdf(out, pdf)
        print(f"pdf: {pdf} — {pdf.stat().st_size / 1e6:.0f} МБ")
    if todo:
        print("\nзаполнить руками:")
        for item in dict.fromkeys(todo):
            print("  ·", item)
    return 0


if __name__ == "__main__":
    sys.exit(main())
