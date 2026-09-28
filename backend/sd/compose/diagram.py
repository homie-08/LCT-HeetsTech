"""Схема процесса из нативных фигур — то, что делает дизайнер, когда в шаблоне
нет макета под шаги.

Карточки со скруглением, номер в кружке, пиктограмма по смыслу шага и текст;
между карточками — стрелки. Больше четырёх шагов — два ряда змейкой, как в
«изогнутом процессе» SmartArt. Всё редактируется в PowerPoint: фигуры, текст,
цвета. Цвета и кегли берутся из дизайн-системы шаблона; ни одного значения
здесь нет — только доли и правила.
"""

from __future__ import annotations

import math

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from ..fit.textmetrics import metrics_for, wrap_lines
from ..icons import add_icon, pick
from ..icons.shape import IconColor, SCHEME_TO_THEME, strip_style
from ..ooxml.color import relative_luminance
from ..template.model import DesignSystem

# Доли карточки: поле от края, диаметр кружка с номером, скругление.
PAD = 0.1
BADGE = 0.24
CORNER = 0.09
# Зазоры между карточками — доли поля схемы.
GAP_X = 0.035
GAP_Y = 0.08
# Карточка не выше своей ширины с небольшим запасом: в один ряд иначе
# получаются столбы.
MAX_ASPECT = 1.15
# Кегль текста шага: от кегля основного текста шаблона вниз до предела читаемости.
MIN_SIZE_PT = 10.0
FALLBACK_SIZE_PT = 16.0
# Плашка карточки — акцент с прозрачностью: на светлом фоне выходит светлый
# оттенок, на тёмном — тёмный, и в обоих случаях это цвет шаблона, а не новый.
CARD_ALPHA = 0.16
# Фон светлый, если его относительная яркость выше этого порога.
LIGHT_PAGE = 0.4


def add_process(slide, box: list[float], steps: list[str], design: DesignSystem,
                width: int, height: int, notes: list[str], number: int) -> None:
    """Рисует шаги процесса в поле `box` (доли слайда)."""
    count = len(steps)
    if count == 0:
        return
    columns = count if count <= 4 else math.ceil(count / 2)
    rows = math.ceil(count / columns)

    box_x, box_y = box[0] * width, box[1] * height
    box_w, box_h = box[2] * width, box[3] * height
    gap_x, gap_y = box_w * GAP_X, box_h * GAP_Y
    card_w = (box_w - gap_x * (columns - 1)) / columns
    card_h = min((box_h - gap_y * (rows - 1)) / rows, card_w * MAX_ASPECT)
    block_h = rows * card_h + (rows - 1) * gap_y
    top = box_y + (box_h - block_h) / 2

    light = _is_light(design, slide)
    accent = contrasting_color(design, slide, "accent_primary")
    # Текст на плашке — тёмный или светлый слот темы по светлоте страницы:
    # роль text_primary у тёмного шаблона серая и на светлой плашке не читается.
    text_color = _text_slot(design, dark=light)
    # Номер стоит на сплошном акценте: светлым, если акцент тёмный, иначе тёмным.
    accent_hex = design.palette.roles.get("accent_primary") or design.palette.theme.get("accent1")
    accent_light = bool(accent_hex) and relative_luminance(accent_hex) > LIGHT_PAGE
    badge_text = _text_slot(design, dark=accent_light)
    style = design.typography.styles.get("body_l1")
    base_size = float(style.size_pt) if style and style.size_pt else FALLBACK_SIZE_PT
    metrics = metrics_for((style.font if style else "") or "")

    # Змейка: чётный ряд слева направо, нечётный — справа налево, чтобы стрелка
    # между рядами шла вниз, а не через весь слайд назад.
    positions: list[tuple[float, float]] = []
    for index in range(count):
        row, column = divmod(index, columns)
        if row % 2 == 1:
            column = columns - 1 - column
        positions.append((box_x + column * (card_w + gap_x), top + row * (card_h + gap_y)))

    pad = card_w * PAD
    badge = min(card_w, card_h) * BADGE
    size_pt = _text_size(steps, card_w - 2 * pad, card_h - 2 * pad - badge - pad * 0.6,
                         metrics, base_size, design)

    for index, (step, (x, y)) in enumerate(zip(steps, positions), start=1):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(int(x)), Emu(int(y)),
                                      Emu(int(card_w)), Emu(int(card_h)))
        card.name = f"Шаг {index}: карточка"
        card.adjustments[0] = CORNER
        strip_style(card)
        _solid(card.fill, accent, alpha=CARD_ALPHA)
        card.line.fill.background()
        card.text_frame.text = ""

        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(int(x + pad)), Emu(int(y + pad)),
                                        Emu(int(badge)), Emu(int(badge)))
        circle.name = f"Шаг {index}: номер"
        strip_style(circle)
        _solid(circle.fill, accent)
        circle.line.fill.background()
        frame = circle.text_frame
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Emu(0)
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = str(index)
        run.font.bold = True
        run.font.size = Pt(max(MIN_SIZE_PT, badge / Emu(Pt(1)) * 0.5))
        _color(run.font.color, badge_text)

        icon = pick(step)
        if icon is not None:
            add_icon(slide, icon, int(x + card_w - pad - badge), int(y + pad), int(badge),
                     accent, label=f"шаг {index}")

        text_top = y + pad + badge + pad * 0.6
        text = slide.shapes.add_textbox(Emu(int(x + pad)), Emu(int(text_top)),
                                        Emu(int(card_w - 2 * pad)),
                                        Emu(int(y + card_h - pad - text_top)))
        text.name = f"Шаг {index}: текст"
        frame = text.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Emu(0)
        lead, tail = _split(step)
        _paragraph(frame.paragraphs[0], lead, size_pt, text_color, bold=bool(tail))
        if tail:
            _paragraph(frame.add_paragraph(), tail, size_pt, text_color, bold=False)

    _arrows(slide, positions, card_w, card_h, gap_x, gap_y, columns, accent)
    notes.append(f"слайд {number}: схема процесса из фигур — {count} шагов")


def _arrows(slide, positions, card_w, card_h, gap_x, gap_y, columns, accent) -> None:
    """Стрелки между соседними шагами: вправо, влево в нечётном ряду, вниз между рядами."""
    arrow_h = card_h * 0.22
    for index in range(1, len(positions)):
        (x0, y0), (x1, y1) = positions[index - 1], positions[index]
        if abs(y1 - y0) < 1:                              # тот же ряд
            width = gap_x * 0.8
            left = (min(x0, x1) + card_w) + (gap_x - width) / 2
            kind = MSO_SHAPE.RIGHT_ARROW if x1 > x0 else MSO_SHAPE.LEFT_ARROW
            shape = slide.shapes.add_shape(kind, Emu(int(left)), Emu(int(y0 + (card_h - arrow_h) / 2)),
                                           Emu(int(width)), Emu(int(arrow_h)))
        else:                                             # переход на следующий ряд
            width = gap_y * 0.8
            arrow_w = arrow_h
            shape = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW,
                                           Emu(int(x0 + (card_w - arrow_w) / 2)),
                                           Emu(int(y0 + card_h + (gap_y - width) / 2)),
                                           Emu(int(arrow_w)), Emu(int(width)))
        shape.name = f"Стрелка {index}"
        strip_style(shape)
        _solid(shape.fill, accent)
        shape.line.fill.background()


def _text_size(steps: list[str], width_emu: float, height_emu: float, metrics,
               base_size: float, design: DesignSystem | None = None) -> float:
    """Кегль, при котором самый длинный шаг помещается в карточку.

    Спускаемся по шкале шаблона, как укладчик: кегль схемы остаётся «его».
    """
    from ..fit.layout import _scale_below

    usable_pt = width_emu / Emu(Pt(1))
    height_pt = height_emu / Emu(Pt(1))
    ladder = [base_size] + ([step for step in _scale_below(design, base_size)
                             if step >= MIN_SIZE_PT] if design is not None else [])
    if ladder[-1] > MIN_SIZE_PT:
        ladder.append(MIN_SIZE_PT)
    for size in ladder:
        fits = True
        for step in steps:
            lead, tail = _split(step)
            lines = len(wrap_lines(lead, usable_pt, metrics, size))
            if tail:
                lines += len(wrap_lines(tail, usable_pt, metrics, size))
            widest = max((metrics.text_width(word, size) for word in step.split()), default=0)
            if lines * size * 1.2 > height_pt or widest > usable_pt:
                fits = False
                break
        if fits:
            return size
    return MIN_SIZE_PT


def _split(step: str) -> tuple[str, str]:
    from ..fit.layout import _split_thesis

    lead, tail = _split_thesis(step)
    return (lead, tail) if tail else (step, "")


def _paragraph(paragraph, text: str, size_pt: float, color: IconColor, bold: bool) -> None:
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    _color(run.font.color, color)


def _is_light(design: DesignSystem, slide=None) -> bool:
    page = background_hex(design, slide) if slide is not None else None
    page = page or design.palette.roles.get("page_bg") or design.palette.theme.get("lt1")
    if not page:
        return True
    return relative_luminance(page) > LIGHT_PAGE


# Проценты DrawingML записаны в тысячных долях процента.
LUM_SCALE = 1000 * 100
# Имена слотов фона и текста в схеме по умолчанию → слоты темы.
CLR_MAP = {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}


def background_hex(design: DesignSystem, slide) -> str | None:
    """Цвет сплошного фона слайда: свой, макета или мастера. None — не сплошной.

    Читаем XML напрямую: свойство `background` в python-pptx при обращении
    дописывает слайду собственный фон, и белый слайд перекрывал бы макет.
    """
    from pptx.oxml.ns import qn

    for owner in (slide, getattr(slide, "slide_layout", None),
                  getattr(getattr(slide, "slide_layout", None), "slide_master", None)):
        if owner is None:
            continue
        common = owner._element.find(qn("p:cSld"))
        background = common.find(qn("p:bg")) if common is not None else None
        if background is None:
            continue
        properties = background.find(qn("p:bgPr"))
        if properties is None:
            return None                                # ссылка на стиль фона — не судим
        solid = properties.find(qn("a:solidFill"))
        if solid is None:
            return None                                # градиент, картинка — не судим
        scheme = solid.find(qn("a:schemeClr"))
        if scheme is not None:
            slot = scheme.get("val", "")
            value = design.palette.theme.get(CLR_MAP.get(slot, slot))
            if not value:
                return None
            modifier = 0.0
            lum_mod = scheme.find(qn("a:lumMod"))
            lum_off = scheme.find(qn("a:lumOff"))
            if lum_off is not None:
                modifier = int(lum_off.get("val", "0")) / LUM_SCALE
            elif lum_mod is not None:
                modifier = int(lum_mod.get("val", str(LUM_SCALE))) / LUM_SCALE - 1.0
            return _shade(value, modifier)
        rgb = solid.find(qn("a:srgbClr"))
        if rgb is not None:
            return f"#{rgb.get('val', '')}"
        return None
    return None


def _shade(value: str, brightness: float) -> str:
    """Оттенок цвета темы: положительная яркость светлит, отрицательная — темнит."""
    if not brightness:
        return value
    from ..ooxml.color import hex_to_rgb, rgb_to_hex

    r, g, b = hex_to_rgb(value)
    if brightness > 0:
        r, g, b = (c + (1 - c) * brightness for c in (r, g, b))
    else:
        r, g, b = (c * (1 + brightness) for c in (r, g, b))
    return rgb_to_hex((r, g, b))


def contrasting_color(design: DesignSystem, slide, preferred: str = "accent_primary") -> IconColor:
    """Цвет роли, если он читается на фоне слайда, иначе цвет текста для этого фона."""
    from ..ooxml.color import contrast_ratio

    page = background_hex(design, slide)
    wanted = design.palette.roles.get(preferred) or design.palette.theme.get("accent1")
    if page and wanted and contrast_ratio(wanted, page) < MIN_ICON_CONTRAST:
        return _role_color(design, "text_primary" if relative_luminance(page) > LIGHT_PAGE
                           else "text_inverse")
    return _role_color(design, preferred)


# Ниже этого контраста акцент на фоне слайда не читается — берём цвет текста.
MIN_ICON_CONTRAST = 2.0


def _text_slot(design: DesignSystem, dark: bool) -> IconColor:
    """Тёмный (dk1) или светлый (lt1) текст темы; без темы — роли дизайн-системы."""
    slot = "dk1" if dark else "lt1"
    if design.palette.theme.get(slot):
        return ("scheme", slot)
    return _role_color(design, "text_primary" if dark else "text_inverse")


def _role_color(design: DesignSystem, role: str) -> IconColor:
    """Цвет роли дизайн-системы: слот темы, если роль из него взята, иначе значение."""
    origin = design.palette.role_provenance.get(role, "")
    if origin.startswith("theme."):
        slot = origin.split(".", 1)[1]
        if slot in SCHEME_TO_THEME:
            return ("scheme", slot)
    value = design.palette.roles.get(role)
    # Роль, выведенная из употребления, часто совпадает со слотом темы —
    # тогда фигура перекрасится вместе с темой, как и остальной шаблон.
    for slot, hex_value in design.palette.theme.items():
        if value and hex_value and hex_value.upper() == value.upper() and slot in SCHEME_TO_THEME:
            return ("scheme", slot)
    if not value:
        fallback = {"accent_primary": "accent1", "text_primary": "dk1",
                    "text_inverse": "lt1"}[role]
        return ("scheme", fallback)
    return ("rgb", value)


def _solid(fill, color: IconColor, alpha: float | None = None) -> None:
    """Сплошная заливка цветом роли; `alpha` — непрозрачность 0..1."""
    from pptx.oxml.ns import qn

    fill.solid()
    _color(fill.fore_color, color)
    if alpha is not None:
        solid = fill._xPr.find(qn("a:solidFill"))
        node = solid[0] if solid is not None and len(solid) else None
        if node is not None:
            for old in node.findall(qn("a:alpha")):
                node.remove(old)
            alpha_node = node.makeelement(qn("a:alpha"), {"val": str(int(alpha * LUM_SCALE))})
            node.append(alpha_node)


def _color(color_format, color: IconColor) -> None:
    from pptx.dml.color import RGBColor

    kind, value = color
    if kind == "scheme":
        color_format.theme_color = SCHEME_TO_THEME[value]
    else:
        color_format.rgb = RGBColor.from_string(value.lstrip("#"))
