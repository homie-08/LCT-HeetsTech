"""HTTP-интерфейс сервиса.

Ручки повторяют CLI: разобрать шаблон, собрать презентацию, посмотреть отчёт.
Тяжёлые шаги (рендер, сборка) уходят в фоновую задачу — интерфейс опрашивает
состояние и показывает прогресс, а не ждёт ответа полминуты.
"""

from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

from fastapi import BackgroundTasks, FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse

from ..content import parse_content, parse_text
from ..llm import build_client, llm_status
from ..plan import build_plan, match_deck
from ..qa import build_with_repair
from ..render.preview import render_patterns
from ..template.analyze import analyze_template
from .storage import Job, store

TEMPLATE_SUFFIXES = {".pptx", ".potx"}

app = FastAPI(title="Цифровой дизайнер презентаций", version="0.1.0")
app.add_middleware(
    CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])


# --- разбор шаблона ---------------------------------------------------------

def _analyze_template_job(job_id: str, path: Path) -> None:
    job = store.get(job_id)
    if job is None:
        return
    try:
        store.update(job_id, status="running", stage="разбор шаблона")
        design = analyze_template(path)
        (job.dir / "design.json").write_text(
            json.dumps(design.model_dump(mode="json", by_alias=True),
                       ensure_ascii=False, indent=2), encoding="utf-8")

        store.update(job_id, stage="рендер паттернов")
        previews: dict[str, Path] = {}
        try:
            from ..ooxml.package import Package

            previews = render_patterns(design, Package.open(path), job.dir / "previews",
                                       width_px=900)
        except Exception as error:                     # рендерера может не быть
            store.update(job_id, stage=f"рендер недоступен: {error}")

        store.update(
            job_id, status="ready", stage="готово",
            file=path.name,
            source=design.source.model_dump(mode="json"),
            palette=design.palette.model_dump(mode="json"),
            typography=design.typography.model_dump(mode="json"),
            geometry=design.geometry.model_dump(mode="json"),
            warnings=design.warnings,
            patterns=[{
                "id": pattern.id,
                "archetype": pattern.archetype,
                "name": pattern.source_name,
                "mode": pattern.render_mode,
                "confidence": pattern.confidence,
                "evidence": pattern.evidence,
                "slots": [slot.model_dump(mode="json") for slot in pattern.slots],
                "repeat": pattern.repeat.max if pattern.repeat else None,
                "preview": f"/api/templates/{job_id}/preview/{pattern.id}"
                           if pattern.id in previews else None,
            } for pattern in design.patterns],
        )
        for pattern_id, image in previews.items():
            target = job.dir / "previews" / f"{pattern_id}.png"
            target.parent.mkdir(parents=True, exist_ok=True)
            if Path(image).resolve() != target.resolve():
                shutil.copyfile(image, target)
    except Exception as error:                          # pragma: no cover
        store.fail(job_id, str(error))


@app.post("/api/templates")
async def upload_template(background: BackgroundTasks, file: UploadFile = File(...)):
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in TEMPLATE_SUFFIXES:
        raise HTTPException(400, "нужен файл .pptx или .potx")

    job = store.create("template")
    path = job.dir / (file.filename or "template.pptx")
    path.write_bytes(await file.read())
    store.update(job.id, file=path.name)

    background.add_task(_analyze_template_job, job.id, path)
    return {"id": job.id, "status": job.status}


@app.get("/api/templates")
def list_templates():
    return [job.to_json() for job in store.list("template")]


@app.get("/api/templates/{template_id}")
def get_template(template_id: str):
    job = _require(template_id, "template")
    return job.to_json()


@app.get("/api/templates/{template_id}/preview/{pattern_id}")
def pattern_preview(template_id: str, pattern_id: str):
    job = _require(template_id, "template")
    path = job.dir / "previews" / f"{pattern_id}.png"
    if not path.exists():
        raise HTTPException(404, "превью не готово")
    return FileResponse(path, media_type="image/png")


# --- сборка презентации -----------------------------------------------------

def _defect_json(defect) -> dict:
    """Находка аудита для интерфейса: поля модели плюс устойчивый адрес."""
    payload = defect.model_dump(mode="json")
    payload["id"] = defect.id
    return payload


def _slide_cards(design, spec, filled, images: list[str]) -> list[dict]:
    """Карточки слайдов для интерфейса: замысел, макет, заполненность, снимок."""
    cards = []
    for index, (slide, fill) in enumerate(zip(spec.slides, filled)):
        pattern = design.pattern(slide.pattern_id)
        cards.append({
            "n": slide.n,
            "intent": slide.plan.intent,
            "key_message": slide.plan.key_message,
            "pattern": slide.pattern_id,
            "pattern_name": pattern.source_name if pattern else "",
            "archetype": pattern.archetype if pattern else "",
            "why": slide.why,
            "used_slots": fill.used_slots,
            "total_slots": len(fill.fills),
            "image": images[index] if index < len(images) else None,
        })
    return cards


def _render_slides(deck_path: Path, folder: Path, url_prefix: str) -> list[str]:
    from ..render import render_deck

    rendered = render_deck(deck_path, folder, 1100)
    return [f"{url_prefix}/{index}" for index, _ in enumerate(rendered.images, start=1)]


def _export_pptx_variant(job_id: str, entry: dict, folder: Path, title: str) -> None:
    """Экспорт собранного варианта в .pdf и .html; ссылки дописываются в карточку.

    Неудача экспорта — заметка в карточке, а не провал задачи: pptx уже готов.
    """
    from ..render.export import to_html, to_pdf

    deck_path = folder / "deck.pptx"
    images = sorted((folder / "slides").glob("slide-*.png"))
    try:
        to_html(deck_path, images, folder / "deck.html", title=title)
        entry["html"] = f"/api/decks/{job_id}/html?variant={entry['id']}"
    except Exception as error:                      # noqa: BLE001 — в заметки
        entry["notes"].append(f"html недоступен: {error}")
    try:
        to_pdf(deck_path, folder / "deck.pdf")
        entry["pdf"] = f"/api/decks/{job_id}/pdf?variant={entry['id']}"
    except Exception as error:                      # noqa: BLE001 — в заметки
        entry["notes"].append(f"pdf недоступен: {error}")


def _build_deck_job(job_id: str, template_path: Path, content_dir: Path,
                    brief: str, slides: int | None, use_llm: bool) -> None:
    """Три варианта вёрстки одного контента на одном шаблоне.

    Варианты различаются плотностью — сколько мыслей уносит слайд и какие
    раскладки под них берутся; шаблон, контент и правила укладки у всех одни.
    Первым собирается сбалансированный: он же отдаётся как основной.
    """
    from ..pipeline import build_variants
    from ..plan import variants as variant_set

    job = store.get(job_id)
    if job is None:
        return
    try:
        store.update(job_id, status="running", stage="разбор контента")
        ir = parse_content(content_dir)

        store.update(job_id, stage="разбор шаблона")
        design = analyze_template(template_path)

        client = build_client() if use_llm else None
        order = ["balanced", "compact", "spacious"]
        results = build_variants(
            design, ir, template_path, job.dir, client=client, brief=brief,
            target_slides=slides, names=order,
            progress=lambda variant: store.update(
                job_id, stage=f"вариант «{variant.label}»: план, макеты, сборка"))

        variants_json: list[dict] = []
        for name in variant_set.ORDER:
            built = results[name]
            store.update(job_id, stage=f"вариант «{built.variant.label}»: рендер слайдов")
            images: list[str] = []
            try:
                images = _render_slides(built.path, job.dir / name / "slides",
                                        f"/api/decks/{job_id}/slide/{name}")
            except Exception as error:
                built.report.notes.append(f"рендер недоступен: {error}")
            variants_json.append({
                "id": name,
                "label": built.variant.label,
                "tagline": built.variant.tagline,
                "plan_source": built.plan.source,
                "slides": _slide_cards(design, built.spec, built.filled, images),
                "metrics": built.report.metrics.model_dump(mode="json"),
                "defects": [_defect_json(defect) for defect in built.report.defects],
                "checks": built.report.checks,
                "skipped_checks": built.report.skipped_checks,
                "notes": built.report.notes + built.result.notes[:10],
                "mode_counts": built.result.mode_counts,
                "skills": built.report.skills,
                "download": f"/api/decks/{job_id}/file?variant={name}",
                "pdf": None,
                "html": None,
            })

        main = next(item for item in variants_json if item["id"] == variant_set.DEFAULT)
        fields = dict(
            template=design.source.file,
            source=len(ir.blocks),
            plan_source=main["plan_source"],
            variant=variant_set.DEFAULT,
            variants=variants_json,
            slides=main["slides"],
            metrics=main["metrics"],
            defects=main["defects"],
            notes=main["notes"],
            mode_counts=main["mode_counts"],
            download=main["download"],
        )
        # Колода готова — отдаём её сразу; pdf и html доезжают следом, и пока
        # они готовятся, задача помечена `exporting`, чтобы интерфейс дождался.
        store.update(job_id, status="ready", stage="готово", exporting=True, **fields)
        title = ir.meta.title or brief or "Презентация"
        for entry in variants_json:
            _export_pptx_variant(job_id, entry, job.dir / entry["id"], title)
            store.update(job_id, variants=variants_json, pdf=main["pdf"],
                         html=main["html"], notes=main["notes"],
                         exporting=entry is not variants_json[-1])
    except Exception as error:                      # noqa: BLE001 — в статус задачи
        store.update(job_id, status="failed", stage="ошибка", error=str(error))


@app.post("/api/decks")
async def create_deck(background: BackgroundTasks,
                      template_id: str = Form(...),
                      brief: str = Form(""),
                      slides: str = Form(""),
                      use_llm: str = Form("true"),
                      files: list[UploadFile] = File(...)):
    template_job = _require(template_id, "template")
    template_path = template_job.dir / str(template_job.data.get("file", ""))
    if not template_path.exists():
        raise HTTPException(400, "шаблон не найден")

    job = store.create("deck")
    content_dir = job.dir / "content"
    content_dir.mkdir(parents=True, exist_ok=True)
    for item in files:
        name = Path(item.filename or "content.md").name
        (content_dir / name).write_bytes(await item.read())

    store.update(job.id, template_id=template_id, brief=brief,
                 target_slides=int(slides) if slides.strip().isdigit() else None,
                 use_llm=use_llm.lower() not in ("false", "0", "no"),
                 template_path=str(template_path))
    background.add_task(_build_deck_job, job.id, template_path, content_dir, brief,
                        int(slides) if slides.strip().isdigit() else None,
                        use_llm.lower() not in ("false", "0", "no"))
    return {"id": job.id, "status": job.status}


@app.get("/api/decks/{deck_id}")
def get_deck(deck_id: str):
    return _require(deck_id, "deck").to_json()


def _repair_job(job_id: str, variant_id: str, slides: list[int]) -> None:
    """Пересобирает выбранные слайды варианта и обновляет его карточку."""
    from ..pipeline import build_one
    from ..plan import variants as variant_set
    from ..qa import repair_selected

    job = store.get(job_id)
    if job is None:
        return
    data = job.data
    try:
        variant = variant_set.get(variant_id)
        store.update(job_id, status="running",
                     stage=f"ремонт: вариант «{variant.label}», слайды {sorted(slides)}")
        template_path = Path(str(data.get("template_path", "")))
        ir = parse_content(job.dir / "content")
        design = analyze_template(template_path)
        client = build_client() if data.get("use_llm") else None
        plan = build_plan(ir, client, brief=str(data.get("brief", "")),
                          target_slides=data.get("target_slides"), variant=variant)
        spec = match_deck(plan, design, ir, variant)
        folder = job.dir / variant.id
        result, report, filled = repair_selected(
            design, spec, ir, template_path, folder / "deck.pptx", slides,
            client=client, work_dir=folder / "qa")
        report.notes.insert(0, f"вариант вёрстки: {variant.label} — {variant.tagline}")

        images: list[str] = []
        try:
            images = _render_slides(folder / "deck.pptx", folder / "slides",
                                    f"/api/decks/{job_id}/slide/{variant.id}")
        except Exception as error:
            report.notes.append(f"рендер недоступен: {error}")

        variants_json = list(data.get("variants", []))
        entry = {
            "id": variant.id, "label": variant.label, "tagline": variant.tagline,
            "plan_source": plan.source,
            "slides": _slide_cards(design, spec, filled, images),
            "metrics": report.metrics.model_dump(mode="json"),
            "defects": [_defect_json(defect) for defect in report.defects],
            "checks": report.checks, "skipped_checks": report.skipped_checks,
            "notes": report.notes + result.notes[:10],
            "mode_counts": result.mode_counts, "skills": report.skills,
            "download": f"/api/decks/{job_id}/file?variant={variant.id}",
            "repaired": sorted(slides),
            "pdf": None, "html": None,
        }
        _export_pptx_variant(job_id, entry, folder, str(data.get("template", "Презентация")))
        variants_json = [entry if item.get("id") == variant.id else item
                         for item in variants_json]
        fields = dict(status="ready", stage="готово", variants=variants_json)
        if variant.id == data.get("variant", variant_set.DEFAULT):
            fields.update(slides=entry["slides"], metrics=entry["metrics"],
                          defects=entry["defects"], notes=entry["notes"],
                          pdf=entry["pdf"], html=entry["html"])
        store.update(job_id, **fields)
    except Exception as error:                      # noqa: BLE001 — в статус задачи
        store.update(job_id, status="ready", stage=f"ремонт не удался: {error}")


@app.post("/api/decks/{deck_id}/repair")
def repair_deck(deck_id: str, background: BackgroundTasks,
                variant: str = Form("balanced"), defects: str = Form("")):
    """Исправить выбранные находки аудита: пересобрать их слайды ужатым бюджетом.

    `defects` — адреса находок через запятую, как они пришли в отчёте. Чинятся
    только те, у которых `fixable`; по остальным сборке нечего менять — они
    про шаблон или про смысл.
    """
    job = _require(deck_id, "deck")
    if job.status != "ready":
        raise HTTPException(409, "колода ещё собирается")
    if not job.data.get("template_path"):
        raise HTTPException(400, "ремонт по выбору доступен для pptx-шаблонов")
    wanted = {item.strip() for item in defects.split(",") if item.strip()}
    entry = next((item for item in job.data.get("variants", []) if item.get("id") == variant),
                 None)
    if entry is None:
        raise HTTPException(404, "нет такого варианта")
    chosen = [defect for defect in entry.get("defects", [])
              if defect.get("id") in wanted and defect.get("fixable")]
    slides = sorted({int(defect["slide"]) for defect in chosen if defect.get("slide")})
    if not slides:
        raise HTTPException(400, "среди выбранных находок нет исправимых")
    background.add_task(_repair_job, deck_id, variant, slides)
    return {"id": deck_id, "status": "running", "slides": slides}


def _variant_dir(job: Job, variant: str) -> Path:
    """Папка варианта; без варианта — сбалансированный, а для старых задач — корень."""
    from ..plan import variants as variant_set

    name = (variant or variant_set.DEFAULT).strip().lower()
    if name not in variant_set.VARIANTS:
        raise HTTPException(400, f"нет такого варианта вёрстки: {variant}")
    folder = job.dir / name
    return folder if folder.exists() else job.dir


@app.get("/api/decks/{deck_id}/file")
def download_deck(deck_id: str, variant: str = ""):
    job = _require(deck_id, "deck")
    path = _variant_dir(job, variant) / "deck.pptx"
    if not path.exists():
        raise HTTPException(404, "файл ещё не готов")
    stem = Path(str(job.data.get("template", "deck"))).stem
    suffix = f"-{variant}" if variant else ""
    return FileResponse(
        path, filename=f"{stem}{suffix}.pptx",
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation")


@app.get("/api/decks/{deck_id}/slide/{variant}/{number}")
def deck_variant_slide(deck_id: str, variant: str, number: int):
    job = _require(deck_id, "deck")
    path = _variant_dir(job, variant) / "slides" / f"slide-{number:03d}.png"
    if not path.exists():
        raise HTTPException(404, "слайд не отрендерен")
    return FileResponse(path, media_type="image/png")


@app.get("/api/decks/{deck_id}/slide/{number}")
def deck_slide(deck_id: str, number: int):
    job = _require(deck_id, "deck")
    path = job.dir / "slides" / f"slide-{number:03d}.png"
    if not path.exists():
        raise HTTPException(404, "слайд не отрендерен")
    return FileResponse(path, media_type="image/png")


# --- служебное --------------------------------------------------------------

DEMO_ROOT = Path(__file__).resolve().parents[3]
DEMO_TEMPLATES = DEMO_ROOT / "templates" / "deck"
DEMO_CONTENT = DEMO_ROOT / "data" / "content"


@app.get("/api/demo")
def demo_assets():
    """Что можно показать без загрузки файлов вручную — материал для защиты."""
    templates = ([path.name for path in sorted(DEMO_TEMPLATES.iterdir())
                  if path.suffix.lower() in TEMPLATE_SUFFIXES]
                 if DEMO_TEMPLATES.exists() else [])
    content = ([path.name for path in sorted(DEMO_CONTENT.iterdir()) if path.is_dir()]
               if DEMO_CONTENT.exists() else [])
    return {"templates": templates, "content": content}


COVERS = DEMO_ROOT / "data" / "covers"
HTML_TEMPLATES = DEMO_ROOT / "templates" / "incoming" / "package"


def html_templates() -> list[Path]:
    """Колоды-шаблоны в HTML. Обзорная страница и наши сборки — не шаблоны."""
    if not HTML_TEMPLATES.exists():
        return []
    return sorted(path for path in HTML_TEMPLATES.glob("*.dc.html")
                  if not path.name.startswith(("00 ", "our-", "deck.")))


USER_TEMPLATES = DEMO_ROOT / "templates" / "user"
# Те же 18 стилей, но в PowerPoint — с настоящими слайдами-примерами: их
# разбирает PPTX-конвейер и собирает клонированием композиций дизайнера.
PACK_PPTX = HTML_TEMPLATES / "pptx"


def _kind_of(path: Path) -> str:
    return "html" if path.name.endswith(".dc.html") else "pptx"


def _library() -> list[dict]:
    """Все доступные шаблоны: колоды HTML, заготовки PowerPoint и свои."""
    items: list[dict] = []
    for path in html_templates():
        # «11 Neon Lime.dc.html» -> «Neon Lime»
        label = re.sub(r"^\d+\s+", "", path.name.removesuffix(".dc.html"))
        items.append({"id": f"html:{path.name}", "label": label, "kind": "html",
                      "origin": "library"})
    if PACK_PPTX.exists():
        for path in sorted(PACK_PPTX.glob("*.pptx")):
            label = re.sub(r"^\d+\s+", "", path.stem)
            items.append({"id": f"pack:{path.name}", "label": f"{label} (PPTX)",
                          "kind": "pptx", "origin": "library"})
    if DEMO_TEMPLATES.exists():
        for path in sorted(DEMO_TEMPLATES.iterdir()):
            if path.suffix.lower() in TEMPLATE_SUFFIXES:
                items.append({"id": f"pptx:{path.name}", "label": path.stem,
                              "kind": "pptx", "origin": "library"})
    if USER_TEMPLATES.exists():
        for path in sorted(USER_TEMPLATES.iterdir()):
            if path.suffix.lower() in TEMPLATE_SUFFIXES or path.name.endswith(".dc.html"):
                label = path.name.removesuffix(".dc.html") if path.name.endswith(".dc.html") \
                    else path.stem
                items.append({"id": f"mine:{path.name}", "label": label,
                              "kind": _kind_of(path), "origin": "user"})
    return items


def _template_path(template_id: str) -> tuple[Path, str]:
    """Путь к шаблону по идентификатору из библиотеки."""
    kind, _, name = template_id.partition(":")
    name = Path(name).name
    if kind == "html":
        path = HTML_TEMPLATES / name
    elif kind == "pptx":
        path = DEMO_TEMPLATES / name
    elif kind == "pack":
        path = PACK_PPTX / name
    elif kind == "mine":
        path = USER_TEMPLATES / name
    else:
        raise HTTPException(400, "неизвестный шаблон")
    if not path.exists():
        raise HTTPException(404, "шаблон не найден")
    return path, _kind_of(path) if kind == "mine" else kind


UPLOAD_SUFFIXES = TEMPLATE_SUFFIXES | {".html"}


@app.post("/api/library/upload")
def library_upload(file: UploadFile = File(...)):
    """Свой шаблон: файл падает в отдельную папку и попадает в библиотеку.

    Колоде .dc.html для показа нужны рантайм и движок, лежащие рядом, — они
    докладываются из пакета шаблонов, иначе предпросмотр отрисуется пустым.
    """
    name = Path(file.filename or "").name
    suffix = Path(name).suffix.lower()
    if suffix not in UPLOAD_SUFFIXES:
        raise HTTPException(400, "нужен .pptx, .potx или .dc.html")

    USER_TEMPLATES.mkdir(parents=True, exist_ok=True)
    target = USER_TEMPLATES / name
    target.write_bytes(file.file.read())

    if name.endswith((".dc.html", ".html")):
        for runtime in ("support.js", "deck-stage.js"):
            source = HTML_TEMPLATES / runtime
            if source.exists() and not (USER_TEMPLATES / runtime).exists():
                shutil.copyfile(source, USER_TEMPLATES / runtime)
        ds = HTML_TEMPLATES / "_ds"
        if ds.exists() and not (USER_TEMPLATES / "_ds").exists():
            shutil.copytree(ds, USER_TEMPLATES / "_ds")

    return {"id": f"mine:{name}", "label": name.removesuffix(".dc.html")
            if name.endswith(".dc.html") else Path(name).stem,
            "kind": _kind_of(target), "origin": "user"}


@app.get("/api/library")
def library(background: BackgroundTasks):
    """Библиотека шаблонов с обложками. Обложки готовятся в фоне и кешируются."""
    items = _library()
    pending = 0
    for item in items:
        cover = COVERS / f"{_slug(item['id'])}.png"
        if cover.exists():
            item["cover"] = f"/api/covers/{cover.name}"
        else:
            pending += 1
    if pending:
        background.add_task(_render_library_covers)
    return {"templates": items, "pending": pending}


PREVIEWS = COVERS / "previews"
PREVIEW_SLIDES = 4


@app.get("/api/library/preview")
def library_preview(background: BackgroundTasks, id: str):
    """Несколько слайдов шаблона — то, что показывается в предпросмотре.

    Отрисовывается лениво: рисовать по четыре слайда для всех двадцати семи
    шаблонов заранее — это минуты работы браузера ради того, что, возможно,
    никто не откроет. Обложка уже есть, она показывается сразу, остальное
    догружается.
    """
    _template_path(id)                                  # проверка, что шаблон есть
    slug = _slug(id)
    folder = PREVIEWS / slug
    images = sorted(folder.glob("slide-*.png")) if folder.exists() else []

    cover = COVERS / f"{slug}.png"
    urls = [f"/api/library/preview/{slug}/{path.name}" for path in images]
    if not urls and cover.exists():
        urls = [f"/api/covers/{cover.name}"]

    pending = max(0, PREVIEW_SLIDES - len(images))
    if pending:
        background.add_task(_render_preview, id)
    return {"images": urls, "pending": pending}


@app.get("/api/library/preview/{slug}/{filename}")
def library_preview_file(slug: str, filename: str):
    path = PREVIEWS / Path(slug).name / Path(filename).name
    if not path.exists():
        raise HTTPException(404, "слайд ещё не готов")
    return FileResponse(path, media_type="image/png")


def _render_preview(template_id: str) -> None:
    """Готовит первые слайды шаблона для предпросмотра."""
    slug = _slug(template_id)
    folder = PREVIEWS / slug
    if folder.exists() and len(list(folder.glob("slide-*.png"))) >= PREVIEW_SLIDES:
        return
    folder.mkdir(parents=True, exist_ok=True)
    try:
        path, kind = _template_path(template_id)
        if kind == "html":
            from ..htmldeck.export import to_images

            to_images(path, folder, count=PREVIEW_SLIDES, width=960, height=540)
        else:
            _pptx_preview(path, folder)
    except Exception:                                   # рендерера может не быть
        return


def _pptx_preview(source: Path, folder: Path) -> None:
    """Для заготовки PowerPoint показываем её первые макеты, а не пустые слайды."""
    import io

    from pptx import Presentation

    from ..ooxml.clone import drop_all_slides
    from ..ooxml.package import Package
    from ..render import render_deck

    pkg = Package.open(source)
    if pkg.slides:
        # Настоящие слайды дизайнера показательнее пустых макетов.
        result = render_deck(source, folder / "_work", width_px=960)
        for index, image in enumerate(result.images[:PREVIEW_SLIDES], start=1):
            shutil.copyfile(image, folder / f"slide-{index:03d}.png")
        return

    presentation = Presentation(io.BytesIO(pkg.as_pptx_bytes()))
    drop_all_slides(presentation)
    for layout in list(presentation.slide_layouts)[:PREVIEW_SLIDES]:
        presentation.slides.add_slide(layout)

    work = COVERS / "_work"
    work.mkdir(parents=True, exist_ok=True)
    deck = work / f"{source.stem}-preview.pptx"
    presentation.save(str(deck))

    result = render_deck(deck, work / f"{source.stem}-preview", width_px=960)
    for index, image in enumerate(result.images[:PREVIEW_SLIDES], start=1):
        shutil.copyfile(image, folder / f"slide-{index:03d}.png")


@app.get("/api/covers/{filename}")
def cover_file(filename: str):
    path = COVERS / Path(filename).name
    if not path.exists():
        raise HTTPException(404, "обложка ещё не готова")
    return FileResponse(path, media_type="image/png")


def _slug(template_id: str) -> str:
    return re.sub(r"[^\w-]+", "-", template_id).strip("-")


def _render_library_covers() -> None:
    """Готовит обложки всей библиотеки — по одной картинке на шаблон.

    Обложка берётся из самого шаблона: для колоды HTML это снимок её первого
    слайда, для заготовки PowerPoint — отрисованный титульный макет. Рисовать
    что-то своё было бы враньём: в сетке выбора должно быть видно именно то,
    что шаблон даёт.
    """
    COVERS.mkdir(parents=True, exist_ok=True)
    for item in _library():
        target = COVERS / f"{_slug(item['id'])}.png"
        if target.exists():
            continue
        try:
            path, kind = _template_path(item["id"])
            made = (_html_cover(path, target) if kind == "html"
                    else _pptx_cover(path, target))
            if not made:
                continue
        except Exception:                       # рендерера может не быть вовсе
            continue


def _html_cover(deck: Path, target: Path) -> bool:
    from ..htmldeck.export import to_images

    images = to_images(deck, COVERS / "_work" / deck.stem, count=1, width=960, height=540)
    if not images:
        return False
    shutil.copyfile(images[0], target)
    return True


def _pptx_cover(source: Path, target: Path) -> bool:
    import io

    from pptx import Presentation

    from ..ooxml.clone import drop_all_slides
    from ..ooxml.package import Package
    from ..render import render_deck

    pkg = Package.open(source)
    if pkg.slides:
        # У шаблона есть настоящие слайды — обложкой служит первый из них,
        # а не пустой титульный макет.
        result = render_deck(source, COVERS / "_work" / source.stem, width_px=480)
        if result.images:
            shutil.copyfile(result.images[0], target)
            return True
        return False

    presentation = Presentation(io.BytesIO(pkg.as_pptx_bytes()))
    drop_all_slides(presentation)
    presentation.slides.add_slide(list(presentation.slide_layouts)[0])

    work = COVERS / "_work"
    work.mkdir(parents=True, exist_ok=True)
    deck = work / f"{source.stem}.pptx"
    presentation.save(str(deck))

    result = render_deck(deck, work / source.stem, width_px=480)
    if not result.images:
        return False
    shutil.copyfile(result.images[0], target)
    return True


def _cover_path(name: str) -> Path:
    return COVERS / f"{Path(name).stem}.png"


def _render_covers() -> None:
    """Готовит обложки всех демо-шаблонов — по одной картинке на шаблон.

    Обложка делается из самого шаблона: берётся титульный макет, по нему
    заводится один пустой слайд, и он отрисовывается. Рисовать что-то своё было
    бы враньём — в сетке выбора должно быть видно именно то, что шаблон даёт.

    Всё идёт последовательно и в одном проходе: PowerPoint на машине один, и
    девять параллельных обращений он отклоняет.
    """
    import io

    from pptx import Presentation

    from ..ooxml.clone import drop_all_slides
    from ..ooxml.package import Package
    from ..render import render_deck

    COVERS.mkdir(parents=True, exist_ok=True)
    for source in sorted(DEMO_TEMPLATES.iterdir()):
        if source.suffix.lower() not in TEMPLATE_SUFFIXES:
            continue
        target = _cover_path(source.name)
        if target.exists():
            continue
        try:
            pkg = Package.open(source)
            presentation = Presentation(io.BytesIO(pkg.as_pptx_bytes()))
            drop_all_slides(presentation)
            layouts = list(presentation.slide_layouts)
            # Титульный макет — первый по счёту у любого шаблона Office.
            presentation.slides.add_slide(layouts[0])

            work = COVERS / "_work"
            work.mkdir(parents=True, exist_ok=True)
            deck = work / f"{source.stem}.pptx"
            presentation.save(str(deck))

            result = render_deck(deck, work / source.stem, width_px=480)
            if result.images:
                shutil.copyfile(result.images[0], target)
        except Exception:                       # рендерера может не быть вовсе
            continue


@app.get("/api/demo/covers")
def demo_covers(background: BackgroundTasks):
    """Обложки шаблонов: что готово сейчас и запуск отрисовки остального."""
    names = ([path.name for path in sorted(DEMO_TEMPLATES.iterdir())
              if path.suffix.lower() in TEMPLATE_SUFFIXES]
             if DEMO_TEMPLATES.exists() else [])
    covers = {name: f"/api/demo/covers/{Path(name).stem}.png"
              for name in names if _cover_path(name).exists()}
    if len(covers) < len(names):
        background.add_task(_render_covers)
    return {"covers": covers, "pending": len(names) - len(covers)}


@app.get("/api/demo/covers/{filename}")
def demo_cover(filename: str):
    path = COVERS / Path(filename).name
    if not path.exists():
        raise HTTPException(404, "обложка ещё не готова")
    return FileResponse(path, media_type="image/png")


@app.post("/api/demo/template")
def demo_template(background: BackgroundTasks, name: str = Form(...)):
    if ":" in name:
        # Идентификатор из библиотеки — pack:, mine:, pptx: разрешаются одинаково.
        source, _ = _template_path(name)
    else:
        source = DEMO_TEMPLATES / Path(name).name
        if not source.exists():
            # Свои шаблоны лежат отдельно, но собираются тем же конвейером.
            source = USER_TEMPLATES / Path(name).name
    if not source.exists():
        raise HTTPException(404, "шаблон не найден")

    job = store.create("template")
    path = job.dir / source.name
    shutil.copyfile(source, path)
    store.update(job.id, file=path.name)
    background.add_task(_analyze_template_job, job.id, path)
    return {"id": job.id, "status": job.status}


@app.post("/api/demo/deck")
def demo_deck(background: BackgroundTasks, template_id: str = Form(...),
              name: str = Form(...), brief: str = Form(""), slides: str = Form(""),
              use_llm: str = Form("true")):
    template_job = _require(template_id, "template")
    template_path = template_job.dir / str(template_job.data.get("file", ""))
    source = DEMO_CONTENT / Path(name).name
    if not template_path.exists() or not source.is_dir():
        raise HTTPException(404, "шаблон или контент не найдены")

    job = store.create("deck")
    content_dir = job.dir / "content"
    shutil.copytree(source, content_dir, dirs_exist_ok=True)
    store.update(job.id, template_id=template_id, brief=brief,
                 target_slides=int(slides) if slides.strip().isdigit() else None,
                 use_llm=use_llm.lower() not in ("false", "0", "no"),
                 template_path=str(template_path))
    background.add_task(_build_deck_job, job.id, template_path, content_dir, brief,
                        int(slides) if slides.strip().isdigit() else None,
                        use_llm.lower() not in ("false", "0", "no"))
    return {"id": job.id, "status": job.status}


@app.get("/api/status")
def status():
    from ..render import available_backends

    from ..llm import prompts

    return {"llm": llm_status(),
            "renderers": [backend.name for backend in available_backends()],
            "skills": prompts.versions()}


def _html_variant(job: Job, template, template_path: Path, ir, plan, client,
                  variant) -> dict:
    """Один вариант HTML-колоды: разметка и снимки; экспорт — отдельно."""
    from ..htmldeck.build import build_deck
    from ..htmldeck.export import to_images

    folder = job.dir / variant.id
    folder.mkdir(parents=True, exist_ok=True)
    # Колода кладётся рядом с шаблоном: она подтягивает соседние support.js и
    # deck-stage.js, без них страница не откроется.
    deck_path = template_path.parent / f"deck-{job.id}-{variant.id}.dc.html"
    result = build_deck(template, plan, ir, deck_path, client=client)

    images: list[str] = []
    try:
        shots = to_images(deck_path, folder / "slides", len(result.slides))
        images = [f"/api/decks/{job.id}/slide/{variant.id}/{index}"
                  for index, _ in enumerate(shots, start=1)]
    except Exception as error:
        result.warnings.append(f"снимки недоступны: {error}")

    shutil.copyfile(deck_path, folder / "deck.dc.html")
    return {
        "id": variant.id, "label": variant.label, "tagline": variant.tagline,
        "plan_source": plan.source,
        "slides": [{
            "n": slide.n,
            "intent": slide.archetype,
            "key_message": plan.slides[index].key_message if index < len(plan.slides) else "",
            "pattern": slide.label,
            "pattern_name": slide.label,
            "archetype": slide.archetype,
            "why": f"заполнено слотов: {slide.filled}",
            "used_slots": slide.filled,
            "total_slots": slide.filled + sum(1 for note in slide.notes if "пустым" in note),
            "image": images[index] if index < len(images) else None,
        } for index, slide in enumerate(result.slides)],
        "metrics": {"slides": len(result.slides),
                    "filled": sum(slide.filled for slide in result.slides),
                    "rewritten": result.rewritten},
        "notes": result.warnings,
        "labels": [slide.label for slide in result.slides],
        "deck_path": str(deck_path),
        "download": None,
        "pdf": None,
        "html": f"/api/decks/{job.id}/html?variant={variant.id}",
    }


def _export_html_variant(job: Job, entry: dict) -> None:
    """Экспорт одного варианта в pptx и pdf; ссылки дописываются в задачу."""
    from ..htmldeck.export import to_pdf, to_pptx

    deck_path = Path(entry["deck_path"])
    folder = job.dir / entry["id"]
    count = len(entry["slides"])
    try:
        to_pptx(deck_path, folder / "deck.pptx", count, notes=entry["labels"])
        entry["download"] = f"/api/decks/{job.id}/file?variant={entry['id']}"
    except Exception as error:
        entry["notes"].append(f"pptx недоступен: {error}")
    try:
        to_pdf(deck_path, folder / "deck.pdf")
        entry["pdf"] = f"/api/decks/{job.id}/pdf?variant={entry['id']}"
    except Exception as error:
        entry["notes"].append(f"pdf недоступен: {error}")


def _publish_html(job_id: str, template_name: str, entries: list[dict],
                  status: str, stage: str, exporting: bool) -> None:
    """Состояние задачи с тремя вариантами; основные поля — сбалансированный.

    `exporting` говорит интерфейсу, что файлы части вариантов ещё готовятся:
    по нему он продолжает опрашивать задачу — и перестаёт, когда экспорт
    закончен, удался он или нет.
    """
    from ..plan import variants as variant_set

    public = [{key: value for key, value in entry.items()
               if key not in ("deck_path", "labels")} for entry in entries]
    main = next(item for item in public if item["id"] == variant_set.DEFAULT)
    store.update(job_id, status=status, stage=stage, template=template_name,
                 plan_source=main["plan_source"], variant=variant_set.DEFAULT,
                 variants=public, slides=main["slides"], metrics=main["metrics"],
                 notes=main["notes"], download=main["download"], pdf=main["pdf"],
                 html=main["html"], exporting=exporting)


def _build_html_job(job_id: str, template_path: Path, text: str,
                    slides: int | None, use_llm: bool) -> None:
    """Три варианта колоды в HTML-шаблоне: разметка, снимки, pptx и pdf.

    Снимки трёх вариантов готовятся сразу — по ним интерфейс показывает
    колоду. Экспорт в pptx и pdf долгий (браузер измеряет каждый слайд),
    поэтому сначала экспортируется сбалансированный вариант, задача
    объявляется готовой, а два остальных доэкспортируются следом: ссылки
    появятся в карточках вариантов, как только файлы будут.
    """
    from ..htmldeck import parse_template
    from ..plan import variants as variant_set

    job = store.get(job_id)
    if job is None:
        return
    entries: list[dict] = []
    try:
        store.update(job_id, status="running", stage="разбор шаблона")
        template = parse_template(template_path)

        store.update(job_id, stage="разбор контента")
        ir = parse_text(text)
        client = build_client() if use_llm else None

        for name in ("balanced", "compact", "spacious"):
            variant = variant_set.get(name)
            store.update(job_id, stage=f"вариант «{variant.label}»: план и вёрстка")
            plan = build_plan(ir, client, target_slides=slides, variant=variant)
            entries.append(_html_variant(job, template, template_path, ir, plan, client,
                                         variant))
        entries.sort(key=lambda entry: variant_set.ORDER.index(entry["id"]))

        store.update(job_id, stage="экспорт сбалансированного варианта")
        main = next(entry for entry in entries if entry["id"] == variant_set.DEFAULT)
        _export_html_variant(job, main)
        _publish_html(job_id, template.name, entries, "ready", "готово", exporting=True)

        rest = [entry for entry in entries if entry is not main]
        for index, entry in enumerate(rest):
            _export_html_variant(job, entry)
            _publish_html(job_id, template.name, entries, "ready", "готово",
                          exporting=index < len(rest) - 1)
    except Exception as error:                          # pragma: no cover
        store.fail(job_id, str(error))
    finally:
        for entry in entries:
            Path(entry["deck_path"]).unlink(missing_ok=True)


@app.post("/api/decks/html")
def build_html_deck(background: BackgroundTasks, template_id: str = Form(...),
                    text: str = Form(...), slides: str = Form(""),
                    use_llm: str = Form("true")):
    """Собрать презентацию в HTML-шаблоне по тексту запроса."""
    if not text.strip():
        raise HTTPException(400, "нужен текст запроса")
    path, kind = _template_path(template_id)
    if kind != "html":
        raise HTTPException(400, "этот шаблон собирается через /api/decks")

    job = store.create("deck")
    background.add_task(_build_html_job, job.id, path, text,
                        int(slides) if slides.strip().isdigit() else None,
                        use_llm.lower() not in ("false", "0", "no"))
    return {"id": job.id, "status": job.status}


@app.get("/api/decks/{deck_id}/pdf")
def deck_pdf(deck_id: str, variant: str = ""):
    job = _require(deck_id, "deck")
    path = _variant_dir(job, variant) / "deck.pdf"
    if not path.exists():
        raise HTTPException(404, "pdf не готов")
    return FileResponse(path, media_type="application/pdf", filename="презентация.pdf")


@app.get("/api/decks/{deck_id}/html")
def deck_html(deck_id: str, variant: str = ""):
    job = _require(deck_id, "deck")
    folder = _variant_dir(job, variant)
    path = next((candidate for candidate in (folder / "deck.dc.html", folder / "deck.html")
                 if candidate.exists()), None)
    if path is None:
        raise HTTPException(404, "колода не готова")
    return FileResponse(path, media_type="text/html")


@app.delete("/api/decks/{deck_id}")
def delete_deck(deck_id: str):
    """Удаляет проект: файлы колоды, снимки, pdf и запись в индексе."""
    if not store.remove(deck_id, "deck"):
        raise HTTPException(404, "не найдено")
    return {"ok": True}


# Рантайм, который колода подтягивает относительными ссылками рядом с собой.
# Без него страница по ссылке «Поделиться» — просто стопка секций: движок не
# поднимается, слайды не листаются, анимации не играют.
DECK_RUNTIME = {"support.js", "deck-stage.js"}
ASSET_TYPES = {".js": "application/javascript", ".css": "text/css",
               ".json": "application/json", ".md": "text/markdown"}


@app.get("/api/decks/{deck_id}/{asset:path}")
def deck_asset(deck_id: str, asset: str):
    _require(deck_id, "deck")
    root = HTML_TEMPLATES.resolve()
    if asset in DECK_RUNTIME:
        path = root / asset
    elif asset.startswith("_ds/"):
        # Стили дизайн-системы (нужны шаблону Organic) — только изнутри пакета.
        path = (root / asset).resolve()
        if not str(path).startswith(str(root)):
            raise HTTPException(404, "нет такого файла")
    else:
        raise HTTPException(404, "нет такого файла")
    if not path.exists():
        raise HTTPException(404, "нет такого файла")
    return FileResponse(path, media_type=ASSET_TYPES.get(path.suffix.lower(),
                                                         "application/octet-stream"))


@app.post("/api/outline")
def outline(text: str = Form(...), slides: str = Form(""), use_llm: str = Form("true")):
    """Структура будущей презентации: что окажется на каждом слайде.

    Отдельный шаг до сборки — то, что видит пользователь на экране «Структура».
    План строится быстро (без рендера и укладки), поэтому ручка синхронная:
    заводить фоновую задачу ради секунды ожидания незачем.
    """
    if not text.strip():
        raise HTTPException(400, "нужен текст запроса")

    ir = parse_text(text)
    client = build_client() if use_llm.lower() not in ("false", "0", "no") else None
    plan = build_plan(ir, client,
                      target_slides=int(slides) if slides.strip().isdigit() else None)
    return {
        "source": plan.source,
        "title": plan.title,
        "slides": [{"n": item.n, "intent": item.intent,
                    "key_message": item.key_message, "notes": item.notes}
                   for item in plan.slides],
        "warnings": plan.warnings,
    }


@app.get("/api/health")
def health():
    return JSONResponse({"ok": True})


# --- собранный интерфейс ------------------------------------------------------
# Если фронтенд собран, отдаём его отсюда же: демонстрация запускается одной
# командой, без второго терминала с dev-сервером. Монтируется последним, чтобы
# не перехватывать /api/*.

FRONTEND_DIST = Path(__file__).resolve().parents[3] / "frontend" / "dist"


def mount_frontend(application: FastAPI, dist: Path = FRONTEND_DIST) -> bool:
    if not (dist / "index.html").exists():
        return False

    from fastapi.staticfiles import StaticFiles

    class SPAFiles(StaticFiles):
        """Неизвестный путь — это маршрут интерфейса, а не 404."""

        async def get_response(self, path: str, scope):
            response = await super().get_response(path, scope)
            if response.status_code == 404:
                return await super().get_response("index.html", scope)
            return response

    application.mount("/", SPAFiles(directory=str(dist), html=True), name="frontend")
    return True


mount_frontend(app)


def _require(job_id: str, kind: str) -> Job:
    job = store.get(job_id)
    if job is None or job.kind != kind:
        raise HTTPException(404, "не найдено")
    return job
