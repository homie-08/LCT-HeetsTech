"""Сборка презентации: валидность файла, наполнение, два режима.

Главный инвариант здесь — файл должен открываться. Любая висячая ссылка на
связь делает пакет невалидным, и PowerPoint отказывается открывать презентацию
целиком, без подсказки, что именно не так.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pytest
from conftest import ALL_TEMPLATES, PROJECT, requires_templates

from sd.compose import build_deck
from sd.content import parse_content
from sd.fit import fill_deck
from sd.plan import build_plan, match_deck

CONTENT = PROJECT / "data" / "content" / "product"
R_ATTRIBUTE = re.compile(r'r:(?:id|embed|link|dm|lo|qs|cs)="([^"]+)"')
REL_ID = re.compile(r'Id="([^"]+)"')

pytestmark = [requires_templates,
              pytest.mark.skipif(not CONTENT.exists(), reason="нет демо-контента")]


@pytest.fixture(scope="module")
def ir():
    return parse_content(CONTENT)


def _build(template: Path, ir, out_dir: Path):
    from sd.template.analyze import analyze_template

    design = analyze_template(template)
    plan = build_plan(ir, client=None)
    spec = match_deck(plan, design, ir)
    filled = fill_deck(design, spec, ir)
    result = build_deck(design, filled, ir, template, out_dir / f"{template.stem}.pptx")
    return design, filled, result


def dangling_relationships(path: Path) -> dict[str, list[str]]:
    """Ссылки на связи, которых нет в `.rels` части."""
    problems: dict[str, list[str]] = {}
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        for name in names:
            if not re.match(r"ppt/slides/slide\d+\.xml$", name):
                continue
            xml = archive.read(name).decode("utf-8")
            rels_name = name.replace("slides/", "slides/_rels/") + ".rels"
            declared = set(REL_ID.findall(archive.read(rels_name).decode("utf-8"))) \
                if rels_name in names else set()
            missing = sorted(set(R_ATTRIBUTE.findall(xml)) - declared)
            if missing:
                problems[name] = missing
    return problems


def test_deck_is_built(template_path, ir, tmp_path):
    _, filled, result = _build(template_path, ir, tmp_path)
    assert result.path.exists()
    assert result.slides == len(filled) > 3


def test_no_dangling_relationships(template_path, ir, tmp_path):
    """Регресс: клонированные шейпы тянут ссылки на теги, медиа и гиперссылки.

    Оставленный `rId` без описания делает файл невалидным целиком.
    """
    _, _, result = _build(template_path, ir, tmp_path)
    assert dangling_relationships(result.path) == {}


def test_output_opens_as_presentation(template_path, ir, tmp_path):
    from pptx import Presentation

    _, filled, result = _build(template_path, ir, tmp_path)
    presentation = Presentation(str(result.path))
    assert len(presentation.slides) == len(filled)


def test_template_slides_are_replaced_not_appended(template_path, ir, tmp_path):
    """Слайды-примеры шаблона в результат попадать не должны."""
    from pptx import Presentation

    from sd.ooxml.package import Package

    _, filled, result = _build(template_path, ir, tmp_path)
    assert len(Presentation(str(result.path)).slides) == len(filled)
    assert len(Package.open(template_path).slides) >= 0


def test_slides_carry_content(template_path, ir, tmp_path):
    from pptx import Presentation

    _, _, result = _build(template_path, ir, tmp_path)
    presentation = Presentation(str(result.path))
    with_text = [slide for slide in presentation.slides
                 if any(shape.has_text_frame and shape.text_frame.text.strip()
                        for shape in slide.shapes)]
    assert len(with_text) >= len(presentation.slides) * 0.7


def test_charts_and_tables_are_native(template_path, ir, tmp_path):
    """График и таблица вставляются объектами, а не картинкой или текстом."""
    from pptx import Presentation

    design, filled, result = _build(template_path, ir, tmp_path)
    expected_charts = sum(1 for slide in filled for fill in slide.fills
                          if fill.kind == "chart")
    expected_tables = sum(1 for slide in filled for fill in slide.fills
                          if fill.kind == "table")
    if not expected_charts and not expected_tables:
        pytest.skip("в этом шаблоне графики и таблицы не разместились")

    presentation = Presentation(str(result.path))
    charts = sum(1 for slide in presentation.slides for shape in slide.shapes
                 if shape.has_chart)
    tables = sum(1 for slide in presentation.slides for shape in slide.shapes
                 if shape.has_table)
    assert charts == expected_charts
    assert tables == expected_tables


@pytest.mark.skipif(len(ALL_TEMPLATES) < 2, reason="нужно минимум два шаблона")
def test_same_content_differs_across_templates(ir, tmp_path):
    from pptx import Presentation

    shapes_per_template = {}
    for template in ALL_TEMPLATES[:3]:
        _, _, result = _build(template, ir, tmp_path)
        presentation = Presentation(str(result.path))
        shapes_per_template[template.name] = tuple(
            len(slide.shapes) for slide in presentation.slides)
    assert len(set(shapes_per_template.values())) > 1, \
        f"презентации совпали по структуре: {shapes_per_template}"


# --- уборка незаполненных карточек и строк ------------------------------------

def _boxes(shapes, width, height):
    from sd.compose.deck import _shape_box

    return [box for shape in shapes
            if (box := _shape_box(shape, width, height)) is not None]


def test_empty_card_is_removed_whole():
    """Плашка без текста уходит вместе с акцентной полоской, а не остаётся коробкой."""
    from sd.compose.deck import _contains, _drop_empty_cards

    class Slot:
        def __init__(self, bbox, shape_id):
            self.bbox, self.shape_id = bbox, shape_id

    class Element:
        def __init__(self, holder):
            self._holder = holder

        def getparent(self):
            return self._holder

    class Shape:
        def __init__(self, left, top, width, height, holder):
            self.left, self.top = left, top
            self.width, self.height = width, height
            self._element = Element(holder)

    class Holder(list):
        def remove(self, element):
            self.gone.append(element)

    holder = Holder()
    holder.gone = []

    # Карточка 0.1..0.3 по X с надписью внутри и полоской по верхнему краю.
    card = Shape(1000, 1000, 2000, 2000, holder)
    stripe = Shape(1000, 950, 2000, 60, holder)
    label = Shape(1100, 1500, 1800, 400, holder)
    # Соседняя заполненная карточка — её трогать нельзя.
    other = Shape(4000, 1000, 2000, 2000, holder)
    other_label = Shape(4100, 1500, 1800, 400, holder)

    class Slide:
        shapes = [card, stripe, label, other, other_label]

    class Pattern:
        slots = [Slot((0.11, 0.15, 0.18, 0.04), 3),      # пустая надпись
                 Slot((0.41, 0.15, 0.18, 0.04), 4)]      # заполненная

    notes: list[str] = []
    _drop_empty_cards(Slide(), Pattern(), [Pattern.slots[0]], {4},
                      10000, 10000, notes, number=1)

    assert card._element in holder.gone and stripe._element in holder.gone
    assert other._element not in holder.gone, "убрана заполненная карточка"
    assert notes and "карточка" in notes[0]


def test_icon_above_an_empty_card_column_is_removed():
    """Ряд карточек «иконка над текстом»: у карточки без текста уходит и иконка."""
    from sd.compose.deck import _drop_empty_columns

    class Element:
        def __init__(self, holder):
            self._holder = holder

        def getparent(self):
            return self._holder

    class Shape:
        def __init__(self, left, top, width, height, holder, text=""):
            self.left, self.top = left, top
            self.width, self.height = width, height
            self._element = Element(holder)
            self.has_text_frame = bool(text)
            if text:
                self.text_frame = type("Frame", (), {"text": text})()

    class Holder(list):
        def remove(self, element):
            self.gone.append(element)

    holder = Holder()
    holder.gone = []
    # Три колонки: две заполнены, третья (x 0.67..0.91) пуста; над каждой иконка.
    icons = [Shape(int(x * 10000), 3500, 800, 800, holder) for x in (0.05, 0.36, 0.67)]
    filled_text = Shape(500, 5500, 2400, 1800, holder, "тезис")

    class Slide:
        shapes = [*icons, filled_text]

    class Slot:
        def __init__(self, bbox):
            self.bbox = bbox

    empty = Slot([0.67, 0.55, 0.24, 0.18])
    filled_boxes = [[0.05, 0.55, 0.24, 0.18], [0.36, 0.55, 0.24, 0.18]]
    notes: list[str] = []
    _drop_empty_columns(Slide(), [empty], filled_boxes, 10000, 10000, notes, number=4)

    assert holder.gone == [icons[2]._element]
    assert notes == ["слайд 4: иконка пустой карточки убрана (1 фигур)"]


def test_paragraph_keeps_one_properties_element():
    """Два `a:pPr` в абзаце — файл, который PowerPoint не открывает.

    Шаблоны, собранные чужим экспортом, такое содержат; мы копируем абзац как
    образец, и ошибка расходилась по всей колоде. Регресс: после подстановки
    текста свойства абзаца ровно одни и стоят первыми.
    """
    from lxml import etree
    from pptx.oxml.ns import qn

    from sd.ooxml.text import set_text

    body = etree.fromstring(
        '<p:txBody xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
        ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:bodyPr/><a:lstStyle/>'
        '<a:p>'
        '<a:pPr algn="l"><a:lnSpc><a:spcPct val="79167"/></a:lnSpc></a:pPr>'
        '<a:pPr algn="l"><a:lnSpc><a:spcPct val="79167"/></a:lnSpc></a:pPr>'
        '<a:r><a:rPr lang="ru-RU" sz="1800"/><a:t>Образец</a:t></a:r>'
        '</a:p></p:txBody>')

    set_text(body, ["Первая строка", "Вторая строка"])

    paragraphs = body.findall(qn("a:p"))
    assert len(paragraphs) == 2
    for paragraph in paragraphs:
        properties = paragraph.findall(qn("a:pPr"))
        assert len(properties) == 1, "свойства абзаца должны остаться одни"
        assert paragraph[0] is properties[0], "и стоять первыми"
        # Оформление образца при этом сохраняется.
        assert properties[0].get("algn") == "l"
        assert properties[0].find(qn("a:lnSpc")) is not None
    texts = [paragraph.find(".//" + qn("a:t")).text for paragraph in paragraphs]
    assert texts == ["Первая строка", "Вторая строка"]
