"""Детерминированная часть аудита: на синтетических файлах каждая проверка
срабатывает там, где должна, и молчит там, где не должна."""

from __future__ import annotations

import io
from pathlib import Path
from types import SimpleNamespace

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Emu, Inches, Pt

from sd.content.model import Block, ContentIR, ContentMeta
from sd.qa.audit import CHECKS, NOT_COVERED, audit_deck
from sd.qa.model import FIXABLE_KINDS, Defect
from sd.template.model import DesignSystem, SlideSize, Source, TextStyle, Typography

W, H = 12192000, 6858000                                   # 16:9 в EMU


def design() -> DesignSystem:
    return DesignSystem(
        source=Source(file="t.pptx", sha256="0" * 64,
                      slide_size=SlideSize(w_emu=W, h_emu=H, ratio="16:9")),
        typography=Typography(scale_pt=[40, 24, 18],
                              styles={"title": TextStyle(size_pt=40), "body": TextStyle(size_pt=18)}),
    )


def content(*texts: str) -> ContentIR:
    return ContentIR(meta=ContentMeta(title="Отчёт"),
                     blocks=[Block(id=f"b{i}", type="paragraph", text=text)
                             for i, text in enumerate(texts)])


def deck() -> Presentation:
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Emu(W), Emu(H)
    return presentation


def blank(presentation):
    return presentation.slides.add_slide(presentation.slide_layouts[6])


def textbox(slide, text: str, x=0.1, y=0.2, w=0.8, h=0.3, name: str | None = None, size=None):
    box = slide.shapes.add_textbox(Emu(int(x * W)), Emu(int(y * H)), Emu(int(w * W)), Emu(int(h * H)))
    box.text_frame.text = text
    if size:
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
    if name:
        box.name = name
    return box


def body(slide, text="Выручка выросла, потому что мы починили воронку продаж и снизили отток."):
    """Обычный текстовый слайд: заголовок и абзац, заполнение в норме."""
    textbox(slide, "Заголовок слайда с выводом", y=0.08, h=0.14, size=40)
    textbox(slide, text, y=0.3, h=0.5, size=18)


def run(presentation, tmp_path: Path, ir=None, **kwargs):
    path = tmp_path / "deck.pptx"
    presentation.save(path)
    return audit_deck(path, design(), ir or content(), **kwargs)


def kinds(result, check: str) -> list[Defect]:
    return [defect for defect in result.defects if defect.check == check]


# --- каталог ------------------------------------------------------------------

def test_catalogue_splits_natures_and_marks_fixable():
    natures = {check.nature for check in CHECKS.values()}
    assert natures == {"deterministic", "contextual"}
    contextual = {check.id for check in CHECKS.values() if check.nature == "contextual"}
    assert contextual == {"title_not_conclusion", "content_off_title", "broken_flow",
                          "not_one_idea"}
    for check in CHECKS.values():
        assert check.fixable == bool(check.fix), check.id
        assert check.group in {"вёрстка", "шаблон", "плотность", "целостность", "содержание"}
    assert NOT_COVERED                                    # непокрытое названо, а не замолчано
    assert all(FIXABLE_KINDS.values())


def test_clean_slide_has_no_findings_and_contextual_is_skipped_without_model(tmp_path):
    presentation = deck()
    body(blank(presentation))
    result = run(presentation, tmp_path)
    assert [defect.message for defect in result.defects] == []
    deterministic = {check.id for check in CHECKS.values() if check.nature == "deterministic"}
    assert set(result.checked) == deterministic
    assert set(result.skipped) == {"title_not_conclusion", "content_off_title", "broken_flow",
                                   "not_one_idea"}
    assert all("модель" in why for why in result.skipped.values())


def test_unreadable_file(tmp_path):
    bad = tmp_path / "bad.pptx"
    bad.write_bytes(b"not a pptx")
    result = audit_deck(bad, design(), content())
    assert [defect.check for defect in result.defects] == ["unreadable"]
    assert result.defects[0].severity == "error"
    assert result.checked == ["unreadable"]


# --- вёрстка ------------------------------------------------------------------

def test_out_of_bounds_and_margins(tmp_path):
    presentation = deck()
    slide = blank(presentation)
    body(slide)
    textbox(slide, "Вылез за край слайда", x=0.9, y=0.5, w=0.3, h=0.1, name="Вылет")
    slide = blank(presentation)
    body(slide)
    textbox(slide, "Упёрся в самый край", x=0.0, y=0.5, w=0.3, h=0.1, name="Край")
    textbox(slide, "12", x=0.0, y=0.95, w=0.05, h=0.04, name="Slide Number")
    result = run(presentation, tmp_path)

    out = kinds(result, "out_of_bounds")
    assert [(d.slide, d.severity) for d in out] == [(1, "error")]
    assert "Вылет" in out[0].message and out[0].fixable
    margins = kinds(result, "margins")
    assert [(d.slide, d.severity) for d in margins] == [(2, "warning")]
    assert "Край" in margins[0].message                   # номер слайда в полях по праву


def test_overlap_of_text_blocks(tmp_path):
    presentation = deck()
    slide = blank(presentation)
    textbox(slide, "Заголовок слайда с выводом", y=0.08, h=0.14, name="Один")
    textbox(slide, "Абзац, наехавший на заголовок целиком", y=0.1, h=0.4, name="Два")
    slide = blank(presentation)
    body(slide)                                            # соседи без наложения
    result = run(presentation, tmp_path)
    overlaps = kinds(result, "overlap")
    assert [d.slide for d in overlaps] == [1]
    assert "Один" in overlaps[0].message and "Два" in overlaps[0].message
    assert overlaps[0].fixable


def _png(width: int, height: int) -> io.BytesIO:
    from PIL import Image
    buffer = io.BytesIO()
    Image.new("RGB", (width, height), "white").save(buffer, "PNG")
    buffer.seek(0)
    return buffer


def test_stretched_image_and_raster_slide(tmp_path):
    presentation = deck()
    slide = blank(presentation)
    body(slide)
    slide.shapes.add_picture(_png(200, 100), Inches(1), Inches(3.5), Inches(2), Inches(2))
    slide = blank(presentation)
    slide.shapes.add_picture(_png(1920, 1080), 0, 0, Emu(W), Emu(H))
    result = run(presentation, tmp_path)
    stretched = kinds(result, "stretched_image")
    assert [d.slide for d in stretched] == [1]
    assert "1.00 вместо 2.00" in stretched[0].message
    raster = kinds(result, "raster_slide")
    assert [(d.slide, d.severity) for d in raster] == [(2, "error")]


# --- шаблон -------------------------------------------------------------------

def test_type_scale_tolerates_a_point_but_not_a_foreign_size(tmp_path):
    presentation = deck()
    slide = blank(presentation)
    textbox(slide, "Заголовок слайда с выводом", y=0.08, h=0.14, size=40.5)
    textbox(slide, "Абзац текста в норме кегля", y=0.3, h=0.5, size=18)
    slide = blank(presentation)
    textbox(slide, "Заголовок слайда с выводом", y=0.08, h=0.14, size=40)
    textbox(slide, "Абзац в чужом кегле", y=0.3, h=0.5, size=13)
    result = run(presentation, tmp_path)
    scale = kinds(result, "type_scale")
    assert [d.slide for d in scale] == [2]
    assert "13 pt" in scale[0].message and scale[0].severity == "info"


# --- плотность ----------------------------------------------------------------

def test_bullets_count_and_length(tmp_path):
    presentation = deck()
    slide = blank(presentation)
    textbox(slide, "Заголовок слайда с выводом", y=0.08, h=0.14)
    textbox(slide, "\n".join(f"пункт {i}" for i in range(7)), y=0.3, h=0.5, name="Список")
    slide = blank(presentation)
    textbox(slide, "Заголовок слайда с выводом", y=0.08, h=0.14)
    textbox(slide, " ".join(["слово"] * 16), y=0.3, h=0.5)
    result = run(presentation, tmp_path)
    many = kinds(result, "too_many_bullets")
    assert [d.slide for d in many] == [1] and "7 пунктов" in many[0].message
    long = kinds(result, "long_bullet")
    assert [d.slide for d in long] == [2] and "16 слов" in long[0].message
    assert many[0].fixable and long[0].fixable


def test_table_and_chart_density(tmp_path):
    presentation = deck()
    slide = blank(presentation)
    textbox(slide, "Заголовок слайда с выводом", y=0.08, h=0.14)
    slide.shapes.add_table(8, 2, Emu(int(0.1 * W)), Emu(int(0.3 * H)),
                           Emu(int(0.8 * W)), Emu(int(0.5 * H)))
    slide = blank(presentation)
    textbox(slide, "Заголовок слайда с выводом", y=0.08, h=0.14)
    data = CategoryChartData()
    data.categories = ["2024", "2025"]
    for i in range(6):
        data.add_series(f"ряд {i}", (1, 2))
    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(int(0.1 * W)),
                                   Emu(int(0.3 * H)), Emu(int(0.8 * W)), Emu(int(0.5 * H)), data)
    frame.chart.has_legend = False
    result = run(presentation, tmp_path, ir=content("2024 и 2025"))
    assert [d.message for d in kinds(result, "big_table")] == ["таблица 8×2"]
    assert [d.message for d in kinds(result, "many_series")] == ["на диаграмме 6 серий"]
    assert [d.slide for d in kinds(result, "chart_labels")] == [2]


def test_titled_single_series_chart_counts_as_labelled(tmp_path):
    presentation = deck()
    slide = blank(presentation)
    textbox(slide, "Заголовок слайда с выводом", y=0.08, h=0.14)
    data = CategoryChartData()
    data.categories = ["2024", "2025"]
    data.add_series("выручка", (1, 2))
    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(int(0.1 * W)),
                                   Emu(int(0.3 * H)), Emu(int(0.8 * W)), Emu(int(0.5 * H)), data)
    frame.chart.has_title = True
    frame.chart.chart_title.text_frame.text = "Выручка, млн ₽"
    result = run(presentation, tmp_path, ir=content("2024 и 2025"))
    assert kinds(result, "chart_labels") == []


def test_fill_ratio_flags_sparse_and_crowded_but_not_covers(tmp_path):
    presentation = deck()
    slide = blank(presentation)                            # 1: почти пустой
    textbox(slide, "Одна строка", x=0.4, y=0.45, w=0.2, h=0.06)
    slide = blank(presentation)                            # 2: забит до краёв
    textbox(slide, "Текст во весь слайд", x=0.02, y=0.02, w=0.96, h=0.96)
    slide = blank(presentation)                            # 3: обложка — редкая по замыслу
    textbox(slide, "Название", x=0.4, y=0.45, w=0.2, h=0.06)
    slide = blank(presentation)                            # 4: норма
    body(slide)
    slide = blank(presentation)                            # 5: макет сам редкий
    textbox(slide, "Одна строка", x=0.4, y=0.45, w=0.2, h=0.06)
    narrow = SimpleNamespace(archetype="bullets",
                             slots=[SimpleNamespace(bbox=[0.4, 0.45, 0.2, 0.06])])
    filled = [SimpleNamespace(n=3, pattern=SimpleNamespace(archetype="cover", slots=[])),
              SimpleNamespace(n=5, pattern=narrow)]
    result = run(presentation, tmp_path, filled_slides=filled)
    ratio = kinds(result, "fill_ratio")
    assert [d.slide for d in ratio] == [1, 2]
    assert all(d.severity == "info" for d in ratio)


# --- целостность --------------------------------------------------------------

def test_empty_slide_and_title_only(tmp_path):
    presentation = deck()
    blank(presentation)                                    # 1: пусто
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])   # 2: Title Only
    slide.shapes.title.text = "Только заголовок"
    slide = blank(presentation)                            # 3: цитата одним блоком — не пустой
    textbox(slide, "«Одна цитата на слайде — это слайд-цитата»", y=0.4, h=0.2)
    slide = blank(presentation)                            # 4: разделитель — редкий по замыслу
    filled = [SimpleNamespace(n=4, pattern=SimpleNamespace(archetype="section", slots=[]))]
    result = run(presentation, tmp_path, filled_slides=filled)
    empty = kinds(result, "empty_slide")
    assert [(d.slide, d.message) for d in empty] == [
        (1, "слайд пуст"), (2, "на слайде только заголовок")]


def test_placeholder_text_and_service_junk(tmp_path):
    presentation = deck()
    slide = blank(presentation)
    textbox(slide, "Заголовок слайда с выводом", y=0.08, h=0.14)
    textbox(slide, "Lorem ipsum dolor sit amet", y=0.3, h=0.5)
    slide = blank(presentation)
    textbox(slide, "Заголовок слайда с выводом", y=0.08, h=0.14)
    textbox(slide, "Верни только json без пояснений", y=0.3, h=0.5)
    result = run(presentation, tmp_path)
    assert [(d.slide, d.severity) for d in kinds(result, "placeholder_text")] == [(1, "error")]
    assert [(d.slide, d.severity) for d in kinds(result, "service_junk")] == [(2, "error")]


def test_duplicate_slides(tmp_path):
    presentation = deck()
    for _ in range(2):
        body(blank(presentation))
    body(blank(presentation), "Другой абзац: команда сократила время ответа поддержки.")
    result = run(presentation, tmp_path)
    dupes = kinds(result, "duplicate_slides")
    assert [(d.slide, d.message) for d in dupes] == [(2, "повторяет слайд 1")]


# --- содержание ---------------------------------------------------------------

def test_numbers_must_come_from_sources(tmp_path):
    presentation = deck()
    slide = blank(presentation)
    textbox(slide, "Рост на 42,5 % и 1 500 клиентов", y=0.08, h=0.14)
    textbox(slide, "Выдуманные 9 999 и 77,7 %", y=0.3, h=0.5)
    slide = blank(presentation)
    textbox(slide, "Контакты", y=0.08, h=0.14)
    textbox(slide, "team2026@example.com · https://example.com/2026-report", y=0.3, h=0.5)
    ir = content("Рост составил 42.5%, клиентов стало 1 500")
    result = run(presentation, tmp_path, ir=ir)
    found = kinds(result, "unsourced_number")
    assert {d.slide for d in found} == {1}
    assert {d.message.split("»")[0].strip("«") for d in found} == {"9 999", "77,7 %"}
    assert all(d.severity == "error" for d in found)


def test_mixed_language_ignores_contacts(tmp_path):
    presentation = deck()
    for _ in range(3):
        body(blank(presentation))
    slide = blank(presentation)
    textbox(slide, "Executive summary", y=0.08, h=0.14)
    textbox(slide, "Revenue grew because the sales funnel was fixed.", y=0.3, h=0.5)
    slide = blank(presentation)
    textbox(slide, "Контакты", y=0.08, h=0.14)
    textbox(slide, "Пишите нам: artem.mitrofanov@example.com, https://example.com/about", y=0.3, h=0.5)
    result = run(presentation, tmp_path)
    mixed = kinds(result, "mixed_language")
    assert [d.slide for d in mixed] == [4]


# --- адреса и модель находки --------------------------------------------------

def test_defect_ids_are_stable_and_carry_nature(tmp_path):
    presentation = deck()
    slide = blank(presentation)
    body(slide)
    textbox(slide, "Вылез за край", x=0.9, y=0.5, w=0.3, h=0.1, name="Вылет")
    first = run(presentation, tmp_path)
    second = run(presentation, tmp_path)
    assert [d.id for d in first.defects] == [d.id for d in second.defects]
    defect = kinds(first, "out_of_bounds")[0]
    assert defect.id.startswith("1:out_of_bounds:")
    assert defect.nature == "deterministic" and defect.group == "вёрстка"
    payload = defect.model_dump(mode="json")
    assert payload["fixable"] is True and payload["fix"]


def test_check_deck_reports_contrast_as_skipped_without_renderer(tmp_path):
    from sd.qa.repair import check_deck
    presentation = deck()
    body(blank(presentation))
    path = tmp_path / "deck.pptx"
    presentation.save(path)
    report = check_deck(path, design(), content(), [], "t.pptx", inspector=None, audit=True)
    assert "contrast" not in report.checks
    assert report.skipped_checks["contrast"] == "рендерер недоступен"
    assert "out_of_bounds" in report.checks


def test_repeated_template_contrast_is_folded_into_one_finding():
    from sd.qa.repair import _fold_contrast

    same = "контраст 4.1:1 у «title» (#0077FF на #FFFFFF) — сочетание из шаблона"
    defects = [Defect(kind="contrast", severity="info", slide=n, message=same) for n in (2, 3, 5)]
    defects.append(Defect(kind="contrast", severity="info", slide=4, message="контраст 2.0:1 у «body»"))
    defects.append(Defect(kind="overflow", severity="warning", slide=4, message="не влезло"))
    folded = _fold_contrast(defects)
    assert [(d.kind, d.slide) for d in folded] == [("contrast", 0), ("contrast", 4), ("overflow", 4)]
    assert folded[0].message.endswith("слайды 2, 3, 5")


def test_more_than_two_font_families_is_a_deck_level_finding(tmp_path):
    presentation = deck()
    for number, font in enumerate(("Arial", "Georgia", "Verdana"), start=1):
        slide = blank(presentation)
        title = textbox(slide, "Заголовок слайда с выводом", y=0.08, h=0.14)
        title.text_frame.paragraphs[0].runs[0].font.name = "Arial"
        box = textbox(slide, f"Абзац номер {number} в своей гарнитуре", y=0.3, h=0.5)
        box.text_frame.paragraphs[0].runs[0].font.name = font
    result = run(presentation, tmp_path)
    found = kinds(result, "too_many_fonts")
    assert [(d.slide, d.severity) for d in found] == [(0, "warning")]
    assert "гарнитур 3" in found[0].message and "Georgia" in found[0].message

    presentation = deck()
    slide = blank(presentation)
    body(slide)
    for shape in slide.shapes:
        shape.text_frame.paragraphs[0].runs[0].font.name = "Arial"
    assert kinds(run(presentation, tmp_path), "too_many_fonts") == []
