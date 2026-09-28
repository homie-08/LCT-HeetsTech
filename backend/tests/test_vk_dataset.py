"""Сквозная сборка на трёх шаблонах из датасета организаторов.

Это не проверка красоты — её делает глаз, — а проверка того, что нельзя
терять: таблица приходит нативной, а не строками через «|»; обложка первая;
цифры показателя стоят рядом со своей подписью; на слайдах нет заглушек и
текста шаблона; текст остаётся текстом. Пропускается, если датасета нет.
"""

from __future__ import annotations

import time
from pathlib import Path

import pytest
from pptx import Presentation

from sd.content import parse_content
from sd.plan import build_plan, match_deck
from sd.qa import build_with_repair
from sd.template.analyze import analyze_template

DATASET = Path(__file__).resolve().parents[2] / "ресурсы" / "Датасет"
CONTENT = Path(__file__).resolve().parents[2] / "data" / "content" / "product"
TEMPLATES = ["VK Tech шаблон.pptx",
             "VK_WorkSpace_Клиентская_конференция_Шаблон_03.pptx",
             "Шаблон презентации VK Education.pptx"]

pytestmark = pytest.mark.skipif(not DATASET.exists() or not CONTENT.exists(),
                                reason="нет датасета организаторов или демо-контента")


@pytest.fixture(scope="module", params=TEMPLATES, ids=lambda name: name.split(".")[0][:12])
def built(request, tmp_path_factory):
    template = DATASET / request.param
    ir = parse_content(CONTENT)
    design = analyze_template(template)
    spec = match_deck(build_plan(ir, None), design, ir)
    out = tmp_path_factory.mktemp("vk") / "deck.pptx"
    started = time.time()
    result, report, filled = build_with_repair(design, spec, ir, template, out,
                                               rounds=1, visual=False)
    return design, spec, filled, Presentation(str(out)), time.time() - started


def _texts(slide) -> list[str]:
    return [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]


def test_cover_first_and_contacts_last(built):
    design, spec, _, pres, _ = built
    assert design.pattern(spec.slides[0].pattern_id).archetype == "cover"
    # Финальной раскладки у шаблона может не быть — тогда контакты идут на
    # обычный текстовый макет, но они обязаны быть на последнем слайде.
    last = list(pres.slides)[-1]
    assert any("potok@example.com" in text for text in _texts(last))
    assert not any(design.pattern(spec.pattern_id).archetype == "cover"
                   for spec in spec.slides[1:-1]), "обложка в середине колоды"


def test_table_is_native_not_pipe_text(built):
    _, _, _, pres, _ = built
    pipes = sum(1 for slide in pres.slides for text in _texts(slide) if " | " in text)
    tables = sum(1 for slide in pres.slides for shape in slide.shapes if shape.has_table)
    assert pipes == 0, "таблица попала на слайд строками через «|»"
    assert tables == 1, "таблица контента должна стать одной нативной таблицей"


def test_chart_is_native(built):
    _, _, _, pres, _ = built
    assert any(shape.has_chart for slide in pres.slides for shape in slide.shapes)


def test_no_slide_is_a_single_picture(built):
    """Слайд единой картинкой по ТЗ не засчитывается."""
    _, _, _, pres, _ = built
    for slide in pres.slides:
        shapes = list(slide.shapes)
        assert not (len(shapes) == 1 and shapes[0].shape_type == 13)


def test_no_template_text_leaks(built):
    _, _, _, pres, _ = built
    leaked = [text for slide in pres.slides for text in _texts(slide)
              if "Заголовок в одну" in text or "Безопасность" in text
              or "Lorem" in text or "Оцените высокий уровень" in text]
    assert not leaked, leaked


def test_metric_values_keep_their_labels(built):
    """«45 млн ₽» стоит рядом с «экономия…», а не с «7 месяцев»."""
    _, _, filled, _, _ = built
    for slide in filled:
        values = {fill.text for fill in slide.fills if fill.text in ("45 млн ₽", "7 месяцев")}
        if {"45 млн ₽", "7 месяцев"} <= values:
            by_text = {fill.text: fill.slot for fill in slide.fills if fill.text}
            labels = {fill.text: fill.slot for fill in slide.fills
                      if fill.text.startswith(("экономия", "окупаемость"))}
            if labels:
                money, payback = by_text["45 млн ₽"], by_text["7 месяцев"]
                for label, slot in labels.items():
                    owner = money if label.startswith("экономия") else payback
                    assert abs(slot.bbox[0] - owner.bbox[0]) < 0.12, f"подпись «{label}» не под своей цифрой"


def test_no_syllable_breaks_expected(built):
    """Слова не шире своих строк: проверка на уровне укладки — по заметкам."""
    _, _, filled, _, _ = built
    for slide in filled:
        for fill in slide.fills:
            assert "не влезает даже на минимальном кегле" not in " ".join(fill.notes) or fill.is_empty


def test_build_is_fast_enough(built):
    """ТЗ: не более пяти минут на колоду. Без модели — секунды."""
    *_, elapsed = built
    assert elapsed < 120


def test_repair_moves_an_overfull_slide_to_a_roomier_layout(tmp_path):
    """Ремонт находки «не размещён фрагмент» — смена макета, а не ужатие."""
    from sd.plan.model import SlidePlan
    from sd.qa import repair_selected

    template = DATASET / TEMPLATES[2]                     # VK Education
    ir = parse_content(CONTENT)
    design = analyze_template(template)
    plan = build_plan(ir, None)
    # Шесть тезисов на слайде — больше, чем мест в кольцевом макете на четыре.
    heavy = next(slide for slide in plan.slides if slide.intent in ("bullets", "process"))
    heavy.items = [f"Тезис номер {number} про обработку документов" for number in range(1, 7)]
    spec = match_deck(plan, design, ir)
    spec_slide = next(item for item in spec.slides if item.n == heavy.n)
    spec_slide.pattern_id = "bullets-slide17"              # кольцо на четыре места

    result, report, filled = repair_selected(design, spec, ir, template, tmp_path / "deck.pptx",
                                             [heavy.n], visual=False)
    repaired = next(slide for slide in filled if slide.n == heavy.n)
    assert repaired.pattern.id != "bullets-slide17"
    assert not repaired.dropped
    assert any("макет заменён" in note for note in report.notes)
