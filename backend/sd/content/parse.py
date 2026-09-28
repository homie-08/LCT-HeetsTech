"""Разбор входного контента в Content IR.

Разбор детерминированный: структуру даёт сам файл, а не модель. LLM подключается
позже и только там, где нужен смысл — план истории и сокращение формулировок.
Это важно для воспроизводимости: один и тот же контент всегда даёт одни и те же
блоки, и расхождения в презентациях объясняются шаблоном, а не разбором.
"""

from __future__ import annotations

import csv
import re
from pathlib import Path

from markdown_it import MarkdownIt

from .model import Asset, Block, ContentIR, ContentMeta

TEXT_SUFFIXES = {".md", ".markdown", ".txt"}
TABLE_SUFFIXES = {".csv", ".tsv", ".xlsx", ".xlsm"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}

# Что читается как метрика: доля, кратность, денежный объём, «в N раз».
METRIC_RE = re.compile(
    r"(?P<value>"
    r"[+−-]?\d+(?:[.,]\d+)?\s?(?:%|×|x|х)"
    r"|[+−-]?\d+(?:[.,]\d+)?\s?(?:млн|млрд|тыс\.?)(?:\s?(?:₽|руб\.?|\$|€))?"
    r"|[+−-]?\d+(?:[.,]\d+)?\s?(?:₽|\$|€)"
    r"|в\s\d+(?:[.,]\d+)?\s?раза?"
    r"|\d+(?:[.,]\d+)?\s(?:рабочих\s)?(?:дн(?:я|ей|ь)|недел\w*|месяц\w*|лет|год\w*)"
    r")", re.IGNORECASE)

BULLET_RE = re.compile(r"^\s*(?:[-*•–—]|\d+[.)])\s+(?P<item>.+)$")
CONTACT_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+|\+?\d[\d\s()-]{8,}\d|https?://\S+")

# Предлоги и союзы, которые остаются висеть на конце подписи, когда число
# вынули из середины фразы.
HANGING_TAIL_RE = re.compile(
    r"\s+(?:до|от|за|на|в|с|по|около|более|менее|свыше|примерно|почти|уже|"
    r"и|или|а|но|же|чем|как)$", re.IGNORECASE)


class _Counter:
    def __init__(self) -> None:
        self.n = 0

    def next(self, prefix: str = "b") -> str:
        self.n += 1
        return f"{prefix}{self.n}"


# --- markdown / текст --------------------------------------------------------

def _markdown_blocks(text: str, ids: _Counter, source: str) -> list[Block]:
    parser = MarkdownIt("commonmark").enable("table")
    tokens = parser.parse(text)
    blocks: list[Block] = []
    index = 0

    while index < len(tokens):
        token = tokens[index]

        if token.type == "heading_open":
            blocks.append(Block(id=ids.next(), type="heading", level=int(token.tag[1:]),
                                text=tokens[index + 1].content.strip(), source=source))
            index += 3
            continue

        if token.type == "paragraph_open":
            content = tokens[index + 1].content.strip()
            if content:
                blocks.append(_paragraph_block(content, ids, source))
            index += 3
            continue

        if token.type in ("bullet_list_open", "ordered_list_open"):
            ordered = token.type.startswith("ordered")
            depth, items = 1, []
            index += 1
            while index < len(tokens) and depth > 0:
                inner = tokens[index]
                if inner.type in ("bullet_list_open", "ordered_list_open"):
                    depth += 1
                elif inner.type in ("bullet_list_close", "ordered_list_close"):
                    depth -= 1
                elif inner.type == "inline":
                    items.append(inner.content.strip())
                index += 1
            if items:
                blocks.append(Block(id=ids.next(), type="steps" if ordered else "list",
                                    items=items, source=source))
            continue

        if token.type == "blockquote_open":
            depth = 1
            quoted: list[str] = []
            index += 1
            while index < len(tokens) and depth > 0:
                inner = tokens[index]
                if inner.type == "blockquote_open":
                    depth += 1
                elif inner.type == "blockquote_close":
                    depth -= 1
                elif inner.type == "inline":
                    quoted.append(inner.content.strip())
                index += 1
            if quoted:
                blocks.append(Block(id=ids.next(), type="quote", text=" ".join(quoted),
                                    source=source))
            continue

        if token.type == "table_open":
            header, rows, current, in_header = [], [], [], False
            index += 1
            while index < len(tokens) and tokens[index].type != "table_close":
                inner = tokens[index]
                if inner.type == "thead_open":
                    in_header = True
                elif inner.type == "thead_close":
                    in_header = False
                elif inner.type == "tr_open":
                    current = []
                elif inner.type == "tr_close":
                    (header.extend(current) if in_header else rows.append(current))
                elif inner.type == "inline":
                    current.append(inner.content.strip())
                index += 1
            if header or rows:
                blocks.append(Block(id=ids.next(), type="table", header=header,
                                    rows=rows, source=source))
            index += 1
            continue

        index += 1
    return blocks


# Строка короче этого и без признаков предложения — скорее заголовок или пункт
# списка, чем абзац.
SHORT_LINE = 90
SHORT_WORDS = 12
# Столько коротких строк подряд — уже перечень, а не набор заголовков.
LIST_RUN = 3


# Точка в конце выдаёт законченную мысль. У короткой строки она ничего не
# значит («Заключение.»), у длинной — это предложение, а не заголовок.
SENTENCE_LINE = 45


def _looks_short(block: Block) -> bool:
    """Похожа ли строка на заголовок или пункт списка, а не на абзац."""
    if block.type != "paragraph" or not block.text:
        return False
    text = block.text.strip()
    if len(text) > SENTENCE_LINE and text.endswith((".", "!", "?")):
        return False
    return len(text) <= SHORT_LINE and len(text.split()) <= SHORT_WORDS


def _structure_if_flat(blocks: list[Block], ids: _Counter, source: str) -> list[Block]:
    """Выводит структуру, если разметки в тексте не оказалось вовсе."""
    if any(block.type in ("heading", "list", "steps", "table") for block in blocks):
        return blocks
    return infer_structure(blocks, ids, source)


def infer_structure(blocks: list[Block], ids: _Counter, source: str) -> list[Block]:
    """Размечает структуру в тексте, набранном без разметки.

    Пользователь вставляет в поле реферат или конспект, а не markdown: заголовки
    там — просто короткие строки, а список — несколько таких строк подряд. Без
    разбора это семьдесят абзацев без единого раздела, и презентация выходит
    пустой. Правила простые и проверяемые глазом:

    - короткая строка перед длинным абзацем — заголовок раздела;
    - короткая строка с двоеточием на конце — заголовок перечня;
    - три и более коротких строк подряд — перечень;
    - первая строка документа — название.
    """
    result: list[Block] = []
    index = 0
    first = True

    while index < len(blocks):
        block = blocks[index]
        if not _looks_short(block):
            result.append(block)
            index += 1
            first = False
            continue

        following = blocks[index + 1] if index + 1 < len(blocks) else None
        heads_a_list = block.text.rstrip().endswith(":")
        before_paragraph = following is not None and not _looks_short(following)

        if first or before_paragraph or heads_a_list:
            level = 1 if first else 2
            result.append(Block(id=ids.next(), type="heading", level=level,
                                text=block.text.rstrip(" :"), source=source))
            index += 1
            first = False
            if not heads_a_list:
                continue

        # Собираем перечень: подряд идущие короткие строки. Последняя перед
        # длинным абзацем в список не входит — это заголовок следующего раздела.
        items: list[str] = []
        while index < len(blocks) and _looks_short(blocks[index]):
            nxt = blocks[index + 1] if index + 1 < len(blocks) else None
            if items and (blocks[index].text.rstrip().endswith(":")
                          or (nxt is not None and not _looks_short(nxt))):
                break
            items.append(blocks[index].text)
            index += 1

        if len(items) >= LIST_RUN or (heads_a_list and items):
            result.append(Block(id=ids.next(), type="list", items=items, source=source))
        else:
            for item in items:
                result.append(Block(id=ids.next(), type="heading", level=2,
                                    text=item.rstrip(" :"), source=source))
        first = False
    return result


def _paragraph_block(content: str, ids: _Counter, source: str) -> Block:
    # Перенос строки внутри абзаца — свойство исходного файла, а не текста:
    # автор завернул строки по 80 символов. Если оставить их как есть, абзац
    # приедет на слайд разорванным посреди фразы.
    content = re.sub(r"\s*\n\s*", " ", content).strip()
    if CONTACT_RE.search(content) and len(content) < 200:
        return Block(id=ids.next(), type="contact", text=content, source=source)
    return Block(id=ids.next(), type="paragraph", text=content, source=source)


def _plain_blocks(text: str, ids: _Counter, source: str) -> list[Block]:
    """Простой текст: абзацы по пустой строке, маркеры — по началу строки."""
    blocks: list[Block] = []
    for chunk in re.split(r"\n\s*\n", text):
        lines = [line for line in chunk.splitlines() if line.strip()]
        if not lines:
            continue
        bullets = [match.group("item") for line in lines
                   if (match := BULLET_RE.match(line))]
        if len(bullets) >= 2 and len(bullets) == len(lines):
            ordered = bool(re.match(r"^\s*\d+[.)]", lines[0]))
            blocks.append(Block(id=ids.next(), type="steps" if ordered else "list",
                                items=bullets, source=source))
            continue
        joined = " ".join(line.strip() for line in lines)
        if len(lines) == 1 and len(joined) < 90 and not joined.endswith("."):
            blocks.append(Block(id=ids.next(), type="heading", level=2, text=joined,
                                source=source))
        else:
            blocks.append(_paragraph_block(joined, ids, source))
    return blocks


# --- docx --------------------------------------------------------------------

def _docx_blocks(path: Path, ids: _Counter) -> list[Block]:
    from docx import Document

    document = Document(str(path))
    blocks: list[Block] = []
    pending: list[str] = []
    source = path.name

    def flush() -> None:
        nonlocal pending
        if len(pending) >= 2:
            blocks.append(Block(id=ids.next(), type="list", items=pending, source=source))
        elif pending:
            blocks.append(_paragraph_block(pending[0], ids, source))
        pending = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        style = (paragraph.style.name or "").lower()
        if not text:
            continue
        if "list" in style:
            pending.append(text)
            continue
        flush()
        if style.startswith("heading") or style in ("title", "заголовок"):
            level = int(re.search(r"\d", style).group()) if re.search(r"\d", style) else 1
            blocks.append(Block(id=ids.next(), type="heading", level=level, text=text,
                                source=source))
        else:
            blocks.append(_paragraph_block(text, ids, source))
    flush()

    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        blocks.append(Block(id=ids.next(), type="table", header=rows[0], rows=rows[1:],
                            source=source))
    return blocks


# --- таблицы и ряды ----------------------------------------------------------

def _as_number(value: str) -> float | None:
    cleaned = value.replace(" ", "").replace(" ", "").replace("%", "").replace(",", ".")
    try:
        return float(cleaned)
    except ValueError:
        return None


def _table_to_block(header: list[str], rows: list[list[str]], ids: _Counter,
                    source: str) -> Block:
    """Таблица с числовыми колонками — это ряд данных для графика, а не таблица."""
    numeric_columns: dict[str, list[float]] = {}
    for column_index in range(1, len(header)):
        values = [_as_number(row[column_index]) if column_index < len(row) else None
                  for row in rows]
        if values and all(value is not None for value in values):
            numeric_columns[header[column_index]] = [float(v) for v in values]  # type: ignore[arg-type]

    if numeric_columns and len(rows) >= 2:
        return Block(id=ids.next(), type="series", source=source,
                     x=[row[0] for row in rows], y=numeric_columns,
                     suggest="line" if len(rows) > 6 else "bar")
    return Block(id=ids.next(), type="table", header=header, rows=rows, source=source)


def _tabular_blocks(path: Path, ids: _Counter) -> list[Block]:
    if path.suffix.lower() in (".csv", ".tsv"):
        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = [row for row in csv.reader(handle, delimiter=delimiter) if any(row)]
        return [_table_to_block(rows[0], rows[1:], ids, path.name)] if len(rows) > 1 else []

    from openpyxl import load_workbook

    workbook = load_workbook(str(path), data_only=True)
    blocks: list[Block] = []
    for sheet in workbook.worksheets:
        rows = [["" if cell is None else str(cell) for cell in row]
                for row in sheet.iter_rows(values_only=True)
                if any(cell is not None for cell in row)]
        if len(rows) > 1:
            blocks.append(_table_to_block(rows[0], rows[1:], ids,
                                          f"{path.name}:{sheet.title}"))
    return blocks


# --- метрики -----------------------------------------------------------------

def extract_metrics(blocks: list[Block], ids: _Counter) -> list[Block]:
    """Достаёт из текста измеримые утверждения — материал для слайда «в цифрах»."""
    metrics: list[Block] = []
    for block in blocks:
        if block.type not in ("paragraph", "list", "steps"):
            continue
        for text in ([block.text] if block.text else block.items):
            match = METRIC_RE.search(text)
            if match is None:
                continue
            value = match.group("value").strip()
            label = (text[:match.start()] + " " + text[match.end():]).strip(" ,.—–:")
            label = re.sub(r"\s+", " ", label)
            # Число вынуто из середины фразы, и на его месте остаётся предлог:
            # «аудит занимает до» вместо «аудит занимает до 2 недель». Подпись
            # с висящим предлогом читается как оборванная, поэтому убираем его.
            label = re.sub(HANGING_TAIL_RE, "", label).strip(" ,.—–:")
            if not label:
                continue
            metrics.append(Block(id=ids.next("m"), type="metric", value=value,
                                 label=label[:80], derived_from=block.id,
                                 source=block.source, importance=0.8))
    return metrics


# --- сборка ------------------------------------------------------------------

def parse_text(text: str, source: str = "запрос") -> ContentIR:
    """Разбирает текст, набранный прямо в интерфейсе.

    Тот же разбор, что и для файла: пользователь вправе писать заголовками,
    списками и абзацами, и структура из них должна извлекаться так же.
    """
    ids = _Counter()
    blocks = _markdown_blocks(text, ids, source)
    blocks = _structure_if_flat(blocks, ids, source)
    blocks.extend(extract_metrics(blocks, ids))
    return _assemble(blocks, [], [source], fallback_title="")


def _assemble(blocks: list[Block], assets: list[Asset], sources: list[str],
              fallback_title: str) -> ContentIR:
    """Собирает Content IR: заголовок, подзаголовок, язык и объём."""
    heading = next((block for block in blocks if block.type == "heading"), None)
    first_paragraph = next((block for block in blocks if block.type == "paragraph"), None)
    words = sum(len(block.text.split()) + sum(len(item.split()) for item in block.items)
                for block in blocks)

    return ContentIR(
        meta=ContentMeta(
            title=heading.text if heading else fallback_title,
            subtitle=summarize(first_paragraph.text) if first_paragraph else "",
            language="ru" if _looks_russian(blocks) else "en",
            words=words, sources=sources),
        blocks=blocks, assets=assets)


def parse_content(target: str | Path | list[str | Path]) -> ContentIR:
    """Файл, список файлов или папка -> Content IR."""
    paths: list[Path] = []
    for item in (target if isinstance(target, list) else [target]):
        path = Path(item)
        paths.extend(sorted(p for p in path.rglob("*") if p.is_file())
                     if path.is_dir() else [path])

    # Повествование идёт первым: текст задаёт структуру, таблицы и картинки
    # к ней прикладываются. Иначе презентация начинается с графика.
    order = {**{suffix: 0 for suffix in TEXT_SUFFIXES}, ".docx": 0,
             **{suffix: 1 for suffix in TABLE_SUFFIXES},
             **{suffix: 2 for suffix in IMAGE_SUFFIXES}}
    paths.sort(key=lambda item: (order.get(item.suffix.lower(), 3), item.name))

    ids = _Counter()
    blocks: list[Block] = []
    assets: list[Asset] = []
    sources: list[str] = []

    for path in paths:
        suffix = path.suffix.lower()
        sources.append(path.name)
        if suffix in TEXT_SUFFIXES:
            text = path.read_text(encoding="utf-8")
            parsed = (_markdown_blocks(text, ids, path.name) if suffix != ".txt"
                      else _plain_blocks(text, ids, path.name))
            blocks.extend(_structure_if_flat(parsed, ids, path.name))
        elif suffix == ".docx":
            blocks.extend(_docx_blocks(path, ids))
        elif suffix in TABLE_SUFFIXES:
            blocks.extend(_tabular_blocks(path, ids))
        elif suffix in IMAGE_SUFFIXES:
            asset = Asset(id=ids.next("img"), path=str(path))
            try:
                from PIL import Image
                with Image.open(path) as image:
                    asset.width, asset.height = image.size
            except Exception:
                pass
            assets.append(asset)
            blocks.append(Block(id=ids.next(), type="image", asset_id=asset.id,
                                text=path.stem.replace("_", " "), source=path.name))

    blocks.extend(extract_metrics(blocks, ids))
    return _assemble(blocks, assets, sources,
                     fallback_title=paths[0].stem if paths else "")


def summarize(text: str, limit: int = 160) -> str:
    """Короткое пояснение для обложки.

    Режем по границе предложения, а в крайнем случае — по слову: обрыв на
    середине слова («Заменяет связку и») выглядит как ошибка вёрстки, хотя
    приходит из разбора.
    """
    sentence = re.split(r"(?<=[.!?])\s+", text.strip())[0] if text else ""
    if len(sentence) <= limit:
        return sentence
    kept = sentence[:limit].rsplit(" ", 1)[0]
    # Обрыв посреди оборота («…возникающее, когда тело…») читается как сбой.
    # Если недалеко позади есть граница части предложения, кончаем на ней:
    # мысль остаётся целой, а многоточие честно говорит, что было продолжение.
    # Точка с запятой делит мысль сильнее запятой: по ней остаётся законченная
    # часть, а не начатое и брошенное придаточное.
    for mark in (";", ","):
        boundary = kept.rfind(mark)
        if boundary >= limit * 0.6:
            kept = kept[:boundary]
            break
    return kept.rstrip(" ,;:—–-") + "…"


def _looks_russian(blocks: list[Block]) -> bool:
    sample = " ".join(b.text for b in blocks[:20])
    cyrillic = sum(1 for char in sample if "а" <= char.lower() <= "я")
    return cyrillic > len(sample) * 0.2
