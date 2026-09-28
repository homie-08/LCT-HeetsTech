"""Экспорт pptx-колоды в .html и .pdf."""

from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.util import Emu, Inches

from sd.render import available_backends
from sd.render.export import to_html, to_pdf


def _deck(path):
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Emu(12192000), Emu(6858000)
    for number in range(1, 3):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        box.text_frame.text = f"Слайд {number}: срок обработки упал до 1,5 дня"
    presentation.save(path)
    return path


def test_html_export_is_self_contained_and_carries_the_text(tmp_path):
    from PIL import Image

    deck = _deck(tmp_path / "deck.pptx")
    images = []
    for number in range(1, 3):
        image = tmp_path / f"slide-{number:03d}.png"
        Image.new("RGB", (32, 18), "white").save(image)
        images.append(image)
    page = to_html(deck, images, tmp_path / "deck.html", title="Итоги")
    text = page.read_text(encoding="utf-8")
    assert text.count("data:image/png;base64,") == 2          # картинки внутри
    assert "Слайд 2: срок обработки упал до 1,5 дня" in text
    assert "<title>Итоги</title>" in text
    assert "http" not in text.split("<body>")[1]               # ничего снаружи


def test_html_export_names_missing_frames(tmp_path):
    deck = _deck(tmp_path / "deck.pptx")
    page = to_html(deck, [tmp_path / "none.png"], tmp_path / "deck.html")
    assert "кадр не отрисован" in page.read_text(encoding="utf-8")


@pytest.mark.skipif(not available_backends(), reason="нет PowerPoint или LibreOffice")
def test_pdf_export_has_a_page_per_slide(tmp_path):
    import fitz

    deck = _deck(tmp_path / "deck.pptx")
    pdf = to_pdf(deck, tmp_path / "deck.pdf")
    with fitz.open(pdf) as document:
        assert len(document) == 2
