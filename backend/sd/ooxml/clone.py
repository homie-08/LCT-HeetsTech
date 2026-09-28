"""Перенос слайда-примера в новую презентацию.

Режим `clone_slide` нужен для шаблонов, где макеты пустые, а весь дизайн
нарисован руками на слайдах. Копируем шейпы вместе с оформлением, переносим
картинки в новый пакет и переназначаем связи — иначе в файле останутся ссылки
на медиа, которых в нём нет, и PowerPoint предложит «восстановить» презентацию.

Графики и таблицы намеренно не копируются: их мы вставляем нативно со своими
данными, а чужие остались бы с цифрами из шаблона.
"""

from __future__ import annotations

import copy
import io

from lxml import etree

from .ns import NS, qn
from .package import Package

COPYABLE = {qn("p:sp"), qn("p:pic"), qn("p:cxnSp"), qn("p:grpSp")}
GRAPHIC_FRAME = qn("p:graphicFrame")


def drop_all_slides(presentation) -> None:
    """Убирает слайды вместе со связями — иначе файл считается повреждённым."""
    slide_list = presentation.slides._sldIdLst
    for slide_id in list(slide_list):
        rid = slide_id.get(f"{{{NS['r']}}}id")
        if rid:
            presentation.part.drop_rel(rid)
        slide_list.remove(slide_id)


def clear_shapes(slide) -> None:
    """Снимает автоматически добавленные плейсхолдеры перед клонированием."""
    tree = slide.shapes._spTree
    for element in list(tree):
        if element.tag in COPYABLE or element.tag == GRAPHIC_FRAME:
            tree.remove(element)


IMAGE_SUFFIXES = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".emf",
                  ".wmf", ".svg", ".webp")
RELATIONSHIP_NS = f"{{{NS['r']}}}"


def _content_type_of(pkg: Package, target: str) -> str:
    """Тип содержимого части — из [Content_Types].xml исходного пакета."""
    path = target if target.startswith("/") else f"/{target.lstrip('/')}"
    try:
        types = etree.fromstring(pkg.blob("[Content_Types].xml"))
    except Exception:
        return "application/octet-stream"

    namespace = {"ct": "http://schemas.openxmlformats.org/package/2006/content-types"}
    for override in types.findall("ct:Override", namespace):
        if (override.get("PartName") or "").lower() == path.lower():
            return override.get("ContentType") or "application/octet-stream"

    suffix = path.rsplit(".", 1)[-1].lower() if "." in path else ""
    for default in types.findall("ct:Default", namespace):
        if (default.get("Extension") or "").lower() == suffix:
            return default.get("ContentType") or "application/octet-stream"
    return "application/octet-stream"


def _carry_image(slide, pkg: Package, target: str) -> str:
    """Переносит картинку в новый пакет и возвращает новый rId.

    python-pptx открывает картинку через PIL, чтобы узнать её размеры, а PIL не
    знает ни EMF, ни WMF, ни SVG — а именно ими в шаблонах нарисован векторный
    декор. Падать из-за этого нельзя: кладём такую картинку в пакет как есть,
    с типом содержимого из исходного шаблона. Размеры PowerPoint возьмёт из
    самого шейпа — они уже записаны в его геометрии.
    """
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.opc.package import Part

    blob = pkg.blob(target)
    try:
        _, rid = slide.part.get_or_add_image_part(io.BytesIO(blob))
        return rid
    except Exception:
        pass

    package = slide.part.package
    suffix = f".{target.rsplit('.', 1)[-1]}" if "." in target else ".bin"
    partname = package.next_partname(f"/ppt/media/image%d{suffix}")
    part = Part(partname, _content_type_of(pkg, target), package, blob)
    return slide.part.relate_to(part, RT.IMAGE)


def _remap_relationships(element: etree._Element, pkg: Package, source_part: str,
                         slide) -> list[str]:
    """Переносит связи скопированного шейпа в новый пакет.

    Ссылка на связь — это не только картинка: в шейпе может быть гиперссылка,
    медиа или диаграмма. Любой `rId`, оставшийся без описания в новом пакете,
    делает файл невалидным — PowerPoint отказывается его открывать. Поэтому
    обходим все атрибуты пространства имён связей, а не только `a:blip`.
    """
    notes: list[str] = []
    source_rels = pkg.rels(source_part)
    cache: dict[str, str] = {}
    orphaned: list[etree._Element] = []

    for node in list(element.iter()):
        if not isinstance(node.tag, str):
            continue
        for attribute in [key for key in node.attrib if key.startswith(RELATIONSHIP_NS)]:
            rid = node.get(attribute)
            if not rid:
                continue
            if rid in cache:
                node.set(attribute, cache[rid])
                continue

            relation = source_rels.get(rid)
            if relation is None:
                _drop_reference(node, attribute, orphaned)
                notes.append(f"связь {rid} отсутствует в шаблоне — ссылка снята")
                continue

            if relation.external:
                new_rid = slide.part.relate_to(relation.target, relation.type,
                                               is_external=True)
            elif relation.target.lower().endswith(IMAGE_SUFFIXES) and pkg.has(relation.target):
                new_rid = _carry_image(slide, pkg, relation.target)
            else:
                _drop_reference(node, attribute, orphaned)
                notes.append(f"связь {rid} ({relation.target}) не переносится — снята")
                continue

            cache[rid] = new_rid
            node.set(attribute, new_rid)

    for node in orphaned:
        parent = node.getparent()
        if parent is not None:
            parent.remove(node)
    return notes


# Элементы, существующие только ради ссылки: без неё они невалидны.
REFERENCE_ONLY = {qn("p:tags"), qn("a:hlinkClick"), qn("a:hlinkHover"),
                  qn("p:oleObj"), qn("a:videoFile"), qn("a:audioFile")}


def _drop_reference(node: etree._Element, attribute: str,
                    orphaned: list[etree._Element]) -> None:
    """Снимает ссылку, а вместе с ней и элемент, который без неё не имеет смысла.

    `<p:tags r:id="rId2"/>` без атрибута — это уже не «нет тега», а нарушение
    схемы: PowerPoint отказывается открывать такой файл целиком.
    """
    del node.attrib[attribute]
    if node.tag in REFERENCE_ONLY and not node.attrib:
        orphaned.append(node)


def copy_slide_shapes(pkg: Package, source_part: str, slide) -> list[str]:
    """Копирует шейпы слайда-примера в готовый слайд. Возвращает замечания."""
    source_tree = pkg.xml(source_part).find(".//p:cSld/p:spTree", NS)
    if source_tree is None:
        return ["в слайде-примере нет дерева шейпов"]

    notes: list[str] = []
    target_tree = slide.shapes._spTree
    for element in source_tree:
        if element.tag == GRAPHIC_FRAME:
            notes.append("график/таблица из примера пропущены — вставим свои данные")
            continue
        if element.tag not in COPYABLE:
            continue
        copied = copy.deepcopy(element)
        notes.extend(_remap_relationships(copied, pkg, source_part, slide))
        target_tree.append(copied)
    return notes


def copy_slide_background(pkg: Package, source_part: str, slide) -> None:
    """Переносит фон слайда-примера на собранный слайд.

    У шаблонов, экспортированных из HTML, фон задан на каждом слайде, а не в
    мастере. Без переноса тёмная обложка становится белой, и белый текст на
    ней исчезает. Схема требует, чтобы `p:bg` стоял первым детём `p:cSld`.
    """
    source_bg = pkg.xml(source_part).find(".//p:cSld/p:bg", NS)
    if source_bg is None:
        return

    csld = slide._element.find(qn("p:cSld"))
    if csld is None:
        return
    old = csld.find(qn("p:bg"))
    if old is not None:
        csld.remove(old)

    copied = copy.deepcopy(source_bg)
    _remap_relationships(copied, pkg, source_part, slide)
    csld.insert(0, copied)


def find_shape(slide, shape_id: int):
    """Шейп по идентификатору из паттерна."""
    for shape in slide.shapes:
        if shape.shape_id == shape_id:
            return shape
    return None
