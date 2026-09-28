"""Редактируемый экспорт: перенос надписей в настоящие текстовые поля.

Снимок отдаёт точную картинку, но править её нельзя, а презентацию перед
выступлением правят почти всегда. Здесь проверяется вторая половина экспорта:
измеренные надписи должны лечь поверх фона тем же кеглем, цветом и на то же
место — и остаться текстом.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from sd.htmldeck.export import _add_text_boxes
from sd.htmldeck.textmap import (COLLECT_JS, HIDE_JS, MEASURE_JS, SlideText, TextBox,
                                 parse_rgb, variant)


def _box(**kwargs) -> TextBox:
    base = dict(text="Трение в спорте", x=0.1, y=0.2, w=0.5, h=0.1, size=0.05,
                line=0.06, family="Onest", weight=800, italic=False,
                color="rgb(238, 244, 228)", align="left", upper=False)
    base.update(kwargs)
    return TextBox(**base)                       # type: ignore[arg-type]


def _slide_with(items: list[TextBox]):
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    added = _add_text_boxes(slide, SlideText(index=0, items=items), presentation)
    return slide, presentation, added


def test_text_stays_text():
    """Надпись приезжает текстовым полем, а не частью картинки."""
    slide, _, added = _slide_with([_box()])
    assert added == 1
    shapes = [shape for shape in slide.shapes if shape.has_text_frame]
    assert len(shapes) == 1
    assert shapes[0].text_frame.text == "Трение в спорте"


def test_position_follows_measured_share():
    """Доли переводятся в размер слайда: тот же угол, та же ширина."""
    slide, presentation, _ = _slide_with([_box(x=0.25, y=0.5, w=0.4, h=0.1)])
    shape = next(shape for shape in slide.shapes if shape.has_text_frame)
    # Рамка чуть шире измеренной — на запас под другой перенос строк.
    assert shape.left == pytest.approx(0.25 * presentation.slide_width, rel=0.02)
    assert shape.top == pytest.approx(0.5 * presentation.slide_height, rel=0.02)
    assert shape.width == pytest.approx(0.4 * presentation.slide_width, rel=0.05)


def test_font_size_scales_to_slide_height():
    """Кегль задан долей высоты слайда: 0,05 от 7,5 дюйма — это 27 пунктов."""
    slide, _, _ = _slide_with([_box(size=0.05)])
    run = next(shape for shape in slide.shapes
               if shape.has_text_frame).text_frame.paragraphs[0].runs[0]
    assert run.font.size.pt == pytest.approx(27.0, abs=0.5)


def test_style_carried_over():
    """Начертание, цвет и выключка переносятся как измерены."""
    slide, _, _ = _slide_with([_box(weight=800, italic=True, align="center",
                                    color="rgb(10, 20, 30)", family="Manrope")])
    run = next(shape for shape in slide.shapes
               if shape.has_text_frame).text_frame.paragraphs[0].runs[0]
    assert run.font.bold and run.font.italic
    assert run.font.name == "Manrope"
    assert str(run.font.color.rgb) == "0A141E"


def test_uppercase_applied_from_style():
    """PowerPoint не знает про text-transform: регистр задаём сами."""
    slide, _, _ = _slide_with([_box(text="итоги", upper=True)])
    shape = next(shape for shape in slide.shapes if shape.has_text_frame)
    assert shape.text_frame.text == "ИТОГИ"


def test_negative_offsets_stay_on_slide():
    """Надпись у самого края не уезжает за границу слайда отрицательным полем."""
    slide, _, _ = _slide_with([_box(x=0.0, y=0.0)])
    shape = next(shape for shape in slide.shapes if shape.has_text_frame)
    assert shape.left >= 0 and shape.top >= 0


def test_parse_rgb_handles_alpha_and_garbage():
    assert parse_rgb("rgb(238, 244, 228)") == (238, 244, 228)
    assert parse_rgb("rgba(10, 20, 30, 0.5)") == (10, 20, 30)
    assert parse_rgb("") == (0, 0, 0)


def test_variant_injects_script(tmp_path: Path):
    """Скрипт добавляется в копию рядом с колодой — там же, где движок."""
    deck = tmp_path / "deck.dc.html"
    deck.write_text("<html><body><section>слайд</section></body></html>",
                    encoding="utf-8")
    probe = variant(deck, "window.__x=1;", "measure")
    assert probe.parent == deck.parent
    assert "window.__x=1;" in probe.read_text(encoding="utf-8")
    assert deck.read_text(encoding="utf-8").count("__x") == 0


def test_variant_survives_missing_body(tmp_path: Path):
    """Обрезанная разметка без </body> всё равно получает скрипт."""
    deck = tmp_path / "deck.dc.html"
    deck.write_text("<section>слайд</section>", encoding="utf-8")
    assert "window.__x=1;" in variant(deck, "window.__x=1;", "bg").read_text(encoding="utf-8")


def test_hiding_keeps_the_card_but_not_the_glyphs():
    """Прячем цвет текста, а не элемент: фон и рамка карточки нужны на снимке."""
    # Отбор надписей общий с измерением и visibility читает — смотрим только на
    # то, что делает само сокрытие.
    hiding = HIDE_JS[len(COLLECT_JS):]
    assert "'color', 'transparent'" in hiding
    assert "visibility" not in hiding


def test_measure_and_hide_share_one_rule():
    """Отбор надписей общий: иначе спрячется не то, что перенесено."""
    assert "window.__sdItems" in MEASURE_JS and "window.__sdItems" in HIDE_JS
