"""Рендер-слой: зонд макетов, бэкенды, кеш."""

from __future__ import annotations

import io

import pytest
from conftest import ALL_TEMPLATES, requires_templates

from sd.ooxml.package import Package
from sd.render import available_backends, get_renderer

pytestmark = requires_templates

needs_renderer = pytest.mark.skipif(
    not available_backends(),
    reason="нет ни PowerPoint, ни LibreOffice",
)


def test_conversion_does_not_break_the_source_package(template_path):
    """Регресс: пересборка архива не должна портить исходный пакет.

    `ZipFile.writestr` переписывает смещения прямо в переданном `ZipInfo`, а он
    принадлежит исходному архиву — после такой записи оригинал не читается.
    """
    package = Package.open(template_path)
    data = package.as_pptx_bytes()
    assert len(data) > 0
    theme = package.theme_of(package.masters[0])
    assert package.blob(theme), "исходный пакет перестал читаться после конвертации"
    assert package.slides is not None


def test_potx_converts_to_a_valid_presentation(template_path):
    from pptx import Presentation

    package = Package.open(template_path)
    presentation = Presentation(io.BytesIO(package.as_pptx_bytes()))
    assert len(presentation.slides) == len(package.slides)


def test_layout_probe_deck_has_one_slide_per_layout(template_path):
    from pptx import Presentation

    from sd.render.preview import layout_probe_deck

    package = Package.open(template_path)
    deck, order = layout_probe_deck(package)
    assert deck.exists()
    assert order == list(package.layouts), "порядок слайдов зонда должен совпасть с макетами"

    probe = Presentation(str(deck))
    assert len(probe.slides) == len(order)
    for slide in probe.slides:
        # Зонд должен быть пустым: только оформление шаблона, никакого текста.
        assert all(not shape.has_text_frame or not shape.text_frame.text.strip()
                   for shape in slide.shapes)


@needs_renderer
def test_renders_every_layout(tmp_path):
    """Каждому макету — своя картинка нужной ширины.

    Размер файла ничего не доказывает: пустой белый макет весит триста байт и
    это правильный результат, поэтому проверяем, что PNG читается и совпадает
    по геометрии со слайдом шаблона.
    """
    from PIL import Image

    from sd.render.preview import render_layouts

    package = Package.open(ALL_TEMPLATES[0])
    images = render_layouts(package, tmp_path, width_px=480)
    assert len(images) == len(package.layouts)

    slide_w, slide_h = package.slide_size
    for path in images.values():
        with Image.open(path) as image:
            assert image.width == 480
            assert abs(image.height - 480 * slide_h / slide_w) <= 2


@needs_renderer
def test_render_is_cached_and_still_delivers(tmp_path):
    """Второй прогон берётся из кеша, но каталог назначения всё равно заполняется.

    Раньше при попадании в кеш в `out_dir` оставались картинки прошлого прогона —
    вызывающий смотрел туда и видел устаревший результат.
    """
    from sd.render import render_deck
    from sd.render.preview import layout_probe_deck

    package = Package.open(ALL_TEMPLATES[0])
    deck, _ = layout_probe_deck(package)

    first = render_deck(deck, tmp_path / "a", 480)
    second = render_deck(deck, tmp_path / "b", 480)

    assert "кеш" in second.note
    assert len(second.images) == len(first.images)
    assert all(path.parent == tmp_path / "b" for path in second.images)
    assert second.images[0].read_bytes() == first.images[0].read_bytes()


def test_backend_selection_prefers_fidelity():
    backends = available_backends()
    if not backends:
        pytest.skip("нет доступных бэкендов")
    assert get_renderer().fidelity == max(backend.fidelity for backend in backends)
    assert get_renderer("не существует") is None
